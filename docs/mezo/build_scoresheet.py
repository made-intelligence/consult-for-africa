#!/usr/bin/env python3
"""Build the Mezo facility assessment score sheet (xlsx) from the criteria.

    python3 docs/mezo/build_scoresheet.py

One workbook per assessment: copy it, fill the yellow cells, read the Summary.
Items mirror mezo-facility-assessment-criteria.md version 2026-10-v1 (18
minimums, 46 scored items across nine domains, three tiers). Change both
together.
"""
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "mezo-facility-scoresheet.xlsx")

MINIMUMS = [
    ("A1", "State registration for private health establishments, for the services provided"),
    ("A2", "Medical director of record with full, current MDCN registration and licence"),
    ("A3", "Pharmacy (PCN) and laboratory (MLSCN) registered, or written arrangement with registered providers"),
    ("A4", "NNRA licence where X-ray or C-arm is used"),
    ("A5", "Public liability insurance; every clinician holds indemnity cover"),
    ("A6", "WHO Surgical Safety Checklist used for every operation and retained"),
    ("A7", "Pulse oximetry, BP and ECG for every anaesthetic or sedation; capnography with any airway device"),
    ("A8", "Trained anaesthesia provider present throughout every GA or regional"),
    ("A9", "Resuscitation trolley with defibrillator, checked daily, at every theatre and recovery area"),
    ("A10", "Oxygen with documented reserve, and suction, at every theatre and recovery bed"),
    ("A11", "Generator with automatic changeover plus battery backup for theatre monitors and lights"),
    ("A12", "Validated steam autoclave with cycle log and chemical and biological indicators"),
    ("A13", "Clinical waste segregated and collected by a licensed contractor"),
    ("A14", "Written deteriorating-patient plan and ICU on site or written transfer agreement"),
    ("A15", "Screened, cross-matched blood on site or from a named blood bank in agreed time"),
    ("A16", "Clinical record with operative and anaesthetic notes; consent records risks discussed"),
    ("A17", "Named person responsible for patient data; NDPA 2023 compliant processing"),
    ("A18", "Current fire safety certificate, exits and fire equipment signed off"),
]

DOMAINS = [
    ("B1", "Leadership and governance", [
        "Clinical governance meeting held and minuted",
        "Incident reporting system in use",
        "Incidents reviewed, with actions followed up",
        "Complaints logged and answered",
        "Policies in date",
    ]),
    ("B2", "Theatre and equipment", [
        "Theatre condition and air handling",
        "Operating table",
        "Theatre lights",
        "Diathermy",
        "Equipment maintenance and calibration records",
        "Instrument sets complete for the procedures sought",
    ]),
    ("B3", "Anaesthesia", [
        "Anaesthetic machine checked before each list",
        "Anaesthetic drugs stocked and in date",
        "Controlled drug register kept",
        "Difficult airway kit",
        "Pre-anaesthetic assessment documented",
        "Post-anaesthetic recovery discharge criteria",
    ]),
    ("B4", "Infection prevention", [
        "Hand hygiene audit",
        "Theatre cleaning schedule followed",
        "Sterile store",
        "Surgical site infection surveillance",
        "Antibiotic prophylaxis policy and timing",
    ]),
    ("B5", "Recovery, ward and escalation", [
        "Recovery staffing",
        "Recovery monitoring",
        "Ward nurse ratios, day and night",
        "Early warning scores in use",
        "High dependency capacity",
        "Transfer drills or tested transfer route",
    ]),
    ("B6", "Patient safety and records", [
        "Consent quality",
        "Checklist completion in the traced cases",
        "Medicines reconciliation",
        "Venous thromboembolism assessment",
        "Discharge summaries",
    ]),
    ("B7", "Staffing and credentialing", [
        "Rotas cover nights and weekends",
        "Staff licences verified",
        "Life support training current",
        "Induction for new staff",
    ]),
    ("B8", "Support services", [
        "Laboratory turnaround",
        "Blood availability",
        "Imaging",
        "Pharmacy stock control",
        "Water supply",
    ]),
    ("B9", "Patient experience and outcomes", [
        "Information given before admission",
        "Pain management",
        "Outcomes recorded",
        "Ready to record Mezo outcome measures",
    ]),
]

TIER_ITEMS = [
    ("T2a", 2, "Ward with 24-hour nursing and an on-call doctor"),
    ("T2b", 2, "Blood available within one hour"),
    ("T3a", 3, "High dependency beds on site"),
    ("T3b", 3, "Intensive care on site or within 30 minutes by written agreement"),
    ("T3c", 3, "Blood on site"),
    ("T3d", 3, "Ultraclean air, or a documented surgical site infection rate, for implant work"),
    ("T3e", 3, "Image intensifier where the approved procedures need it"),
]

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="0A7B6E")
INPUT = PatternFill("solid", fgColor="FFF4CC")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"


def build():
    wb = Workbook()

    # ---- Part A
    a = wb.active
    a.title = "A Minimums"
    header(a, ["Ref", "Minimum requirement", "Met?", "Evidence seen", "Note"], [6, 70, 9, 40, 40])
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    a.add_data_validation(dv)
    for r, (ref, text) in enumerate(MINIMUMS, start=2):
        a.cell(row=r, column=1, value=ref)
        a.cell(row=r, column=2, value=text).alignment = WRAP
        for c in (3, 4, 5):
            a.cell(row=r, column=c).fill = INPUT
            a.cell(row=r, column=c).alignment = WRAP
        dv.add(a.cell(row=r, column=3))
    a_last = len(MINIMUMS) + 1

    # ---- Part B
    b = wb.create_sheet("B Scored")
    header(b, ["Domain", "Item", "Score 0-2", "Evidence seen", "Note"], [30, 52, 10, 40, 40])
    dv2 = DataValidation(type="whole", operator="between", formula1="0", formula2="2", allow_blank=True)
    b.add_data_validation(dv2)
    r = 2
    ranges = []
    for ref, name, items in DOMAINS:
        start = r
        for it in items:
            b.cell(row=r, column=1, value=f"{ref} {name}")
            b.cell(row=r, column=2, value=it).alignment = WRAP
            for c in (3, 4, 5):
                b.cell(row=r, column=c).fill = INPUT
                b.cell(row=r, column=c).alignment = WRAP
            dv2.add(b.cell(row=r, column=3))
            r += 1
        ranges.append((ref, name, start, r - 1, len(items)))

    # ---- Part C
    c = wb.create_sheet("C Tier items")
    header(c, ["Ref", "Tier", "Item", "Met?", "Note"], [6, 6, 60, 9, 40])
    dv3 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    c.add_data_validation(dv3)
    for rr, (ref, tier, text) in enumerate(TIER_ITEMS, start=2):
        c.cell(row=rr, column=1, value=ref)
        c.cell(row=rr, column=2, value=tier)
        c.cell(row=rr, column=3, value=text).alignment = WRAP
        for col in (4, 5):
            c.cell(row=rr, column=col).fill = INPUT
        dv3.add(c.cell(row=rr, column=4))

    # ---- Summary
    s = wb.create_sheet("Summary", 0)
    s.column_dimensions["A"].width = 38
    s.column_dimensions["B"].width = 18
    s.column_dimensions["C"].width = 14
    rows = [
        ("Facility", None),
        ("Address", None),
        ("Assessor", None),
        ("Date of visit", None),
        ("Assessor has an interest in the facility?", None),
    ]
    for i, (k, _) in enumerate(rows, start=1):
        s.cell(row=i, column=1, value=k).font = BOLD
        s.cell(row=i, column=2).fill = INPUT

    s.cell(row=7, column=1, value="PART A: MINIMUMS").font = BOLD
    s.cell(row=8, column=1, value="Met")
    s.cell(row=8, column=2, value=f"=COUNTIF('A Minimums'!C2:C{a_last},\"Yes\")")
    s.cell(row=9, column=1, value="Not met or not yet checked")
    s.cell(row=9, column=2, value=f"={len(MINIMUMS)}-B8")
    s.cell(row=10, column=1, value="Joins the network?").font = BOLD
    s.cell(row=10, column=2, value=f'=IF(B8={len(MINIMUMS)},"Yes","No")')

    s.cell(row=12, column=1, value="PART B: DOMAIN SCORES").font = BOLD
    s.cell(row=12, column=2, value="Score").font = BOLD
    s.cell(row=12, column=3, value="Per cent").font = BOLD
    first = 13
    for i, (ref, name, lo, hi, n) in enumerate(ranges):
        row = first + i
        s.cell(row=row, column=1, value=f"{ref} {name}")
        s.cell(row=row, column=2, value=f"=SUM('B Scored'!C{lo}:C{hi})&\" / {n * 2}\"")
        s.cell(row=row, column=3, value=f"=SUM('B Scored'!C{lo}:C{hi})/{n * 2}")
        s.cell(row=row, column=3).number_format = "0%"
    last = first + len(ranges) - 1
    total_items = sum(x[4] for x in ranges)
    tot = last + 1
    s.cell(row=tot, column=1, value="Overall").font = BOLD
    b_last = 1 + total_items
    s.cell(row=tot, column=2, value=f"=SUM('B Scored'!C2:C{b_last})&\" / {total_items * 2}\"")
    s.cell(row=tot, column=3, value=f"=SUM('B Scored'!C2:C{b_last})/{total_items * 2}")
    s.cell(row=tot, column=3).number_format = "0%"
    s.cell(row=tot + 1, column=1, value="Lowest domain")
    s.cell(row=tot + 1, column=3, value=f"=MIN(C{first}:C{last})")
    s.cell(row=tot + 1, column=3).number_format = "0%"
    s.cell(row=tot + 2, column=1, value="Items not yet scored")
    s.cell(row=tot + 2, column=2, value=f"=COUNTBLANK('B Scored'!C2:C{b_last})")

    o, m = f"C{tot}", f"C{tot + 1}"
    t2 = "COUNTIFS('C Tier items'!B2:B8,2,'C Tier items'!D2:D8,\"Yes\")=2"
    t3 = "COUNTIFS('C Tier items'!B2:B8,3,'C Tier items'!D2:D8,\"Yes\")=5"
    tr = tot + 4
    s.cell(row=tr, column=1, value="SURGICAL TIER").font = Font(bold=True, size=13)
    s.cell(row=tr, column=2, value=(
        f'=IF(B10<>"Yes","None: a minimum is not met",'
        f'IF(B{tot + 2}>0,"Incomplete: score every item",'
        f'IF(AND({o}>=0.8,{m}>=0.7,{t2},{t3}),"Tier 3: Major and implant",'
        f'IF(AND({o}>=0.7,{m}>=0.6,{t2}),"Tier 2: Inpatient",'
        f'IF(AND({o}>=0.6,{m}>=0.5),"Tier 1: Day case","Not approved for surgery")))))'
    ))
    s.cell(row=tr, column=2).font = Font(bold=True, size=13)
    s.cell(row=tr + 1, column=1, value="May host consultations")
    s.cell(row=tr + 1, column=2, value='=IF(B10="Yes","Yes","No")')

    s.cell(row=tr + 3, column=1, value="Fill the yellow cells on each sheet. Score 0 not in place, 1 in place but not consistent or not documented, 2 documented and consistently followed.").alignment = WRAP
    s.merge_cells(start_row=tr + 3, start_column=1, end_row=tr + 3, end_column=3)
    s.row_dimensions[tr + 3].height = 45

    red = PatternFill("solid", fgColor="F8D7DA")
    a.conditional_formatting.add(f"C2:C{a_last}", CellIsRule(operator="equal", formula=['"No"'], fill=red))
    b.conditional_formatting.add(f"C2:C{b_last}", CellIsRule(operator="equal", formula=["0"], fill=red))

    wb.save(OUT)
    print("  wrote", os.path.relpath(OUT, os.path.dirname(os.path.dirname(HERE))), f"({len(MINIMUMS)} minimums, {total_items} items)")


if __name__ == "__main__":
    build()
