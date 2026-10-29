"""
Consult for Africa: what we do for private hospitals.

The leave-behind for Ejiro's email blast to a hundred hospitals, which she
follows with calls. It sells the whole firm to a hospital owner or medical
director, from claims and debt recovery to the marketing agency, and ends on
three ways to start. No fees, no client names, no product names: the service is
described, never the platform that delivers it.

The look is the Carbon deck (scripts/build-carbon-receivables-deck.py): white
pages, navy and gold rule at the head, teal kicker, Georgia headlines, surface
boxes with a gold spine and a cream callout.

One content definition drives two renderers, so the team can edit the PPTX and
still have a clean PDF to attach:

    python3 scripts/build-cfa-hospital-services-deck.py

Outputs
    docs/growth/cfa-hospital-services-deck.pdf
    docs/growth/cfa-hospital-services-deck.pptx
    docs/Consult for Africa - Services for Hospitals (Deck).pdf
"""

from __future__ import annotations

import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT_DIR = DOCS / "growth"
STEM = "cfa-hospital-services-deck"
SHARE = DOCS / "Consult for Africa - Services for Hospitals (Deck).pdf"
LOGO = str(DOCS / "c4a-logo.png")
ICON = str(DOCS / "c4a-icon.png")

# A4 landscape in points; both renderers work top-left in points.
PW, PH = 842.0, 595.0
MX = 54.0
W = PW - 2 * MX

NAVY = "#0B3C5D"
GOLD = "#D4AF37"
TEAL = "#1F7A8C"
BODY = "#1F2937"
MUTED = "#6B7280"
SURFACE = "#F1F5F9"
CREAM = "#FBF6E6"
INK = "#182448"
RULE = "#DED9CE"
WHITE = "#FFFFFF"

FOOTER = "Consult for Africa   /   consultforafrica.com   /   Lagos and Abuja"

_SUPP = "/System/Library/Fonts/Supplemental/"


def R(t, size=11, font="sans", color=BODY, sb=0, ls=1.3, track=0):
    """One paragraph. font is sans, sansb, sansi, serif, serifb or serifi."""
    return dict(t=t, size=size, font=font, color=color, sb=sb, ls=ls, track=track)


# =============================================================== renderers ===
class Pdf:
    def __init__(self, path):
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        serif = ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique")
        try:
            pdfmetrics.registerFont(TTFont("Georgia", _SUPP + "Georgia.ttf"))
            pdfmetrics.registerFont(TTFont("Georgia-Bold", _SUPP + "Georgia Bold.ttf"))
            pdfmetrics.registerFont(TTFont("Georgia-Italic", _SUPP + "Georgia Italic.ttf"))
            serif = ("Georgia", "Georgia-Bold", "Georgia-Italic")
        except Exception:
            pass
        sans = ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique")
        try:
            # Embedded, and the same face the PPTX uses, so the two match.
            pdfmetrics.registerFont(TTFont("Arial", _SUPP + "Arial.ttf"))
            pdfmetrics.registerFont(TTFont("Arial-Bold", _SUPP + "Arial Bold.ttf"))
            pdfmetrics.registerFont(TTFont("Arial-Italic", _SUPP + "Arial Italic.ttf"))
            sans = ("Arial", "Arial-Bold", "Arial-Italic")
        except Exception:
            pass
        self.fonts = dict(sans=sans[0], sansb=sans[1], sansi=sans[2],
                          serif=serif[0], serifb=serif[1], serifi=serif[2])
        self.c = canvas.Canvas(str(path), pagesize=(PW, PH))
        self.c.setTitle("Consult for Africa: services for private hospitals")
        self.c.setAuthor("Consult for Africa")
        self.started = False
        self.imgs = {}

    def _hex(self, h):
        from reportlab.lib.colors import HexColor
        return HexColor(h)

    def slide(self):
        if self.started:
            self.c.showPage()
        self.started = True

    def rect(self, x, y, w, h, fill, rounded=False):
        c = self.c
        c.setFillColor(self._hex(fill))
        if rounded:
            c.roundRect(x, PH - y - h, w, h, 8, fill=1, stroke=0)
        else:
            c.rect(x, PH - y - h, w, h, fill=1, stroke=0)

    def line(self, x1, y1, x2, y2, color=RULE, lw=0.9):
        self.c.setStrokeColor(self._hex(color))
        self.c.setLineWidth(lw)
        self.c.line(x1, PH - y1, x2, PH - y2)

    def _wrap(self, text, font, size, maxw, track):
        from reportlab.pdfbase.pdfmetrics import stringWidth

        def width(s):
            return stringWidth(s, font, size) + track * len(s)

        lines, cur = [], ""
        for word in text.split():
            trial = (cur + " " + word).strip()
            if width(trial) <= maxw or not cur:
                cur = trial
            else:
                lines.append(cur)
                cur = word
        lines.append(cur)
        return lines

    def _layout(self, runs, maxw, scale):
        out, total = [], 0.0
        for r in runs:
            size = r["size"] * scale
            font = self.fonts[r["font"]]
            lines = self._wrap(r["t"], font, size, maxw, r["track"])
            lh = size * r["ls"]
            total += r["sb"] * scale + lh * len(lines)
            out.append((r, size, font, lines, lh))
        return out, total

    def text(self, x, y, w, h, runs, align="l", anchor="t"):
        c = self.c
        scale = 1.0
        laid, total = self._layout(runs, w, scale)
        while total > h and scale > 0.8:
            scale -= 0.03
            laid, total = self._layout(runs, w, scale)
        top = y if anchor == "t" else (y + (h - total) / 2 if anchor == "m" else y + h - total)
        cursor = top
        for r, size, font, lines, lh in laid:
            cursor += r["sb"] * scale
            c.setFillColor(self._hex(r["color"]))
            for ln in lines:
                base = PH - (cursor + size * 0.92 + (lh - size) / 2)
                if r["track"]:
                    tx = x
                    c.setFont(font, size)
                    for ch in ln:
                        c.drawString(tx, base, ch)
                        tx += c.stringWidth(ch, font, size) + r["track"]
                else:
                    c.setFont(font, size)
                    if align == "c":
                        c.drawCentredString(x + w / 2, base, ln)
                    elif align == "r":
                        c.drawRightString(x + w, base, ln)
                    else:
                        c.drawString(x, base, ln)
                cursor += lh

    def image(self, path, x, y, h):
        from reportlab.lib.utils import ImageReader
        img = self.imgs.setdefault(path, ImageReader(path))
        iw, ih = img.getSize()
        w = h * iw / ih
        self.c.drawImage(img, x, PH - y - h, width=w, height=h, mask="auto")
        return w

    def image_right(self, path, right, y, h):
        from reportlab.lib.utils import ImageReader
        img = self.imgs.setdefault(path, ImageReader(path))
        iw, ih = img.getSize()
        self.image(path, right - h * iw / ih, y, h)

    def save(self):
        self.c.save()


class Pptx:
    FONT = dict(sans=("Arial", False, False), sansb=("Arial", True, False),
                sansi=("Arial", False, True), serif=("Georgia", False, False),
                serifb=("Georgia", True, False), serifi=("Georgia", False, True))

    def __init__(self, path):
        from pptx import Presentation
        from pptx.util import Pt

        self.path = path
        self.prs = Presentation()
        self.prs.slide_width = Pt(PW)
        self.prs.slide_height = Pt(PH)
        self.prs.core_properties.title = "Consult for Africa: services for private hospitals"
        self.prs.core_properties.author = "Consult for Africa"
        self.blank = self.prs.slide_layouts[6]
        self.s = None

    @staticmethod
    def _rgb(h):
        from pptx.dml.color import RGBColor
        h = h.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    def slide(self):
        self.s = self.prs.slides.add_slide(self.blank)

    def rect(self, x, y, w, h, fill, rounded=False):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Pt
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shp = self.s.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
        if rounded:
            shp.adjustments[0] = min(0.5, 8.0 / max(1.0, min(w, h)))
        shp.fill.solid()
        shp.fill.fore_color.rgb = self._rgb(fill)
        shp.line.fill.background()
        shp.shadow.inherit = False

    def line(self, x1, y1, x2, y2, color=RULE, lw=0.9):
        from pptx.enum.shapes import MSO_CONNECTOR
        from pptx.util import Pt
        ln = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        ln.line.color.rgb = self._rgb(color)
        ln.line.width = Pt(lw)

    def text(self, x, y, w, h, runs, align="l", anchor="t"):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Pt
        box = self.s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, r in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = r["ls"] * 0.86
            if r["sb"]:
                p.space_before = Pt(r["sb"])
            run = p.add_run()
            run.text = r["t"]
            name, bold, italic = self.FONT[r["font"]]
            run.font.name = name
            run.font.bold = bold
            run.font.italic = italic
            run.font.size = Pt(r["size"])
            run.font.color.rgb = self._rgb(r["color"])
            if r["track"]:
                run.font._element.set("spc", str(int(r["track"] * 100)))

    def image(self, path, x, y, h):
        from pptx.util import Pt
        pic = self.s.shapes.add_picture(path, Pt(x), Pt(y), height=Pt(h))
        return pic.width / 12700.0

    def image_right(self, path, right, y, h):
        from pptx.util import Pt
        pic = self.s.shapes.add_picture(path, Pt(0), Pt(y), height=Pt(h))
        pic.left = int(Pt(right)) - pic.width

    def save(self):
        self.prs.save(self.path)


# ============================================================= components ===
class Deck:
    def __init__(self, r, total):
        self.r = r
        self.n = 0
        self.total = total

    def page(self, kicker):
        r = self.r
        r.slide()
        self.n += 1
        r.rect(0, 0, PW, 6, NAVY)
        r.rect(0, 6, PW, 3, GOLD)
        r.rect(MX, 57, 26, 3, GOLD)
        r.text(MX + 36, 51, 500, 16, [R(kicker.upper(), 10.5, "sansb", TEAL)])
        r.image_right(ICON, PW - MX, 48, 30)
        r.rect(MX, PH - 32, 20, 2, GOLD)
        r.text(MX, PH - 27, 500, 12, [R(FOOTER, 8, "sans", MUTED)])
        r.text(PW - MX - 80, PH - 27, 80, 12,
               [R("%02d / %02d" % (self.n, self.total), 8, "sans", MUTED)], align="r")

    def heading(self, text, y=92, size=23, w=W):
        lines = 2 if len(text) * size * 0.5 > w else 1
        h = (size + 8) * lines + 4
        self.r.text(MX, y, w, h, [R(text, size, "serifb", INK, ls=1.25)])
        return y + h + 6

    def lede(self, text, y, w=W, size=12.5):
        self.r.text(MX, y, w, 40, [R(text, size, "serifi", MUTED, ls=1.35)])
        return y + 46

    def kick(self, text, x, y, color=TEAL):
        self.r.text(x, y, 400, 12, [R(text.upper(), 8.5, "sansb", color, track=1.6)])

    def bullets(self, items, x, y, w, size=11, gap=9, row_h=None):
        """items are (lead, body); a gold tick sits beside each."""
        for lead, body in items:
            lines = max(1, int((len(lead) + len(body) + 2) * size * 0.5 / (w - 18)) + 1)
            h = row_h or lines * size * 1.38
            self.r.rect(x, y + size * 0.55, 7, 3, GOLD)
            runs = [R(lead, size, "sansb", NAVY, ls=1.3)] if lead else []
            runs.append(R(body, size, "sans", BODY, ls=1.38, sb=2 if lead else 0))
            if lead:
                h += size * 1.4
            self.r.text(x + 18, y, w - 18, h, runs)
            y += h + gap
        return y

    def card(self, x, y, w, h, head, body, fill=SURFACE, hsize=12, bsize=10.5):
        self.r.rect(x, y, w, h, fill, rounded=True)
        self.r.rect(x, y, 5, h, GOLD)
        self.r.text(x + 18, y + 14, w - 32, h - 24, [
            R(head, hsize, "sansb", NAVY, ls=1.25),
            R(body, bsize, "sans", BODY, sb=4, ls=1.38),
        ])

    def stat(self, x, y, w, h, value, label, vsize=24):
        self.r.rect(x, y, w, h, SURFACE, rounded=True)
        self.r.rect(x, y, 5, h, GOLD)
        self.r.text(x + 18, y + 16, w - 30, vsize + 6, [R(value, vsize, "serifb", NAVY, ls=1.1)])
        self.r.text(x + 18, y + 24 + vsize, w - 30, h - vsize - 30, [R(label, 10, "sans", BODY, ls=1.32)])

    def callout(self, y, text, w=W, x=MX, h=46, fill=CREAM, size=11.5):
        self.r.rect(x, y, w, h, fill, rounded=True)
        self.r.rect(x, y, 5, h, GOLD)
        self.r.text(x + 22, y, w - 40, h, [R(text, size, "sansb", NAVY, ls=1.32)], anchor="m")


# ================================================================ content ===
SERVICES = [
    dict(
        kicker="Claims and debt recovery",
        head="Most of what HMOs owe you is disputed, not refused. Settle the dispute and the claim gets paid.",
        lede="We recover unpaid HMO, corporate and scheme claims claim by claim, with the people at the payor who decide.",
        points=[
            ("Check every claim first", "Against the payor's tariff, the authorisation record and the paperwork. Claims that will not pay as they stand are repaired before anyone chases them."),
            ("Reconcile with each payor", "Match what you billed against what each payor says it owes, and clear the difference with their claims lead."),
            ("Follow up, escalate, settle", "Retrospective authorisations agreed, missing notes resubmitted inside the window, short payments negotiated case by case."),
            ("See every claim", "You see each claim's status from submission to cash, the same view we work from."),
        ],
        signs=["Receivables older than ninety days", "HMOs paying less than you billed", "Nobody owns the reconciliation"],
        terms="The first review is free. After that we are paid from what we recover, so there is no fee on money we do not bring in.",
    ),
    dict(
        kicker="Patients and referrals",
        head="A marketing agency that measures admissions and referrals, not impressions.",
        lede="A specialist healthcare marketing team inside a firm that has run hospitals, so the campaign is built around how patients and doctors actually choose.",
        points=[
            ("Position the hospital", "What you are known for, who it is for and why a patient or a referring doctor should choose you over the hospital down the road."),
            ("Bring patients in", "Search, social, community and corporate campaigns, written with your clinicians and checked for compliance before they run."),
            ("Build the referral base", "A known list of the doctors, clinics and HMOs who send you patients, and a plan to grow it."),
            ("Prove it worked", "A live view of enquiries, bookings and the revenue they became, so spend follows what converts."),
        ],
        signs=["Beds or clinics running below capacity", "Marketing spend you cannot trace to patients", "Referrals concentrated in a few doctors"],
        terms="Monthly retainer, a fixed price project, or both. We agree the numbers it will be judged on before it starts.",
    ),
    dict(
        kicker="Turnaround and financial performance",
        head="When cash tightens, the hospital needs control back before it needs a strategy.",
        lede="We stop the leaks, steady the cash and put a reporting rhythm in place that management can keep after we leave.",
        points=[
            ("Revenue capture", "Billing integrity, price realisation, payer reconciliation and the charges that never reach an invoice."),
            ("Cost discipline", "Procurement, pharmacy and consumable leakage, and staff cost against revenue."),
            ("Productivity", "Theatre utilisation, bed days, clinic throughput and diagnostic yield."),
            ("Cash and control", "A thirteen week cash view, a monthly management pack and a board that sees the same numbers."),
        ],
        signs=["Revenue rising while cash falls", "Salaries are hard to meet", "The monthly numbers arrive late or not at all"],
        terms="A fixed term management mandate or an embedded team against agreed targets.",
    ),
    dict(
        kicker="Organisational audit",
        head="Know what is actually happening in the hospital before you decide what to change.",
        lede="A structured diagnostic across ten domains, from governance and finance to clinical quality and people, delivered as a report the board can act on.",
        points=[
            ("Evidence, not opinion", "Documents, data, site visits, case traces and confidential staff surveys."),
            ("Ten domains", "Governance, strategy, finance, revenue cycle, clinical quality, operations, people, systems, facilities and patient experience."),
            ("A maturity rating", "Where each domain stands today, against what good looks like for a hospital of your size."),
            ("Decisions for the board", "A short list of decisions in order, with what each one costs and returns."),
        ],
        signs=["A new owner, board or chief executive", "Before a raise, a merger or an expansion", "A sense that the numbers and the reality differ"],
        terms="A fixed price project, typically four to eight weeks from site visit to report.",
    ),
]

GROUPED = [
    dict(
        kicker="Strategy and growth",
        terms="A fixed price project with a defined output, and an execution phase if you want us alongside you.",
        head="Grow where the demand is, at a price the market will pay.",
        cards=[
            ("Service lines", "Which services to add, grow or close, from demand, margin and what you are equipped to win."),
            ("Pricing and tariff", "Your tariff against the market, HMO dependency, payer mix and self-pay positioning."),
            ("New sites and feasibility", "Market entry, feasibility studies, business plans and financial models a lender will read."),
            ("Structuring", "Joint ventures, partnerships with visiting specialists and holding company structures, alongside your counsel."),
        ],
    ),
    dict(
        kicker="People and leadership",
        terms="Per placement, a monthly retainer, or a fractional leader on a fixed term.",
        head="Fill the posts, keep the good people, and know who can lead before you appoint them.",
        cards=[
            ("Recruitment and locums", "Permanent clinical and non-clinical hires from a verified pool, and locum cover in days."),
            ("Leadership assessment", "A structured assessment of every senior candidate or sitting leader, with coaching where it helps."),
            ("HR that runs", "Payroll, onboarding, staff recognition and an anonymous staff pulse, set up and handed over."),
            ("Fractional leadership", "An experienced chief operating officer, medical director or commercial lead, two or three days a week."),
        ],
    ),
    dict(
        kicker="Governance and quality",
        terms="A gap analysis first, then a readiness programme against a date you set.",
        head="Quality and governance that protect patients and stand up to a regulator, an insurer or a buyer.",
        cards=[
            ("Clinical governance", "Committees, escalation, incident systems, mortality and morbidity review, and clinical audit."),
            ("Accreditation readiness", "Gap analysis, closure and mock survey for SafeCare, COHSASA and JCI."),
            ("Board governance", "Board and committee set up, delegation of authority and director training."),
            ("Regulatory standing", "Registration, licensing and inspection readiness with the state and federal regulators."),
        ],
    ),
    dict(
        kicker="Systems and specialists",
        terms="Advisory by the project, and specialists arranged session by session.",
        head="The right systems underneath, and the specialists you cannot hire locally.",
        cards=[
            ("Hospital systems", "Choosing an EMR or hospital system, overseeing the rollout and digitising the paper workflows."),
            ("Dashboards", "The few numbers that run the hospital, on one screen, updated without anyone building a spreadsheet."),
            ("Diaspora specialists", "Senior Nigerian consultants abroad for case review, teaching, visiting lists and service set up."),
            ("Health systems", "Programme design for HMOs, state schemes and development partners."),
        ],
    ),
]

TOTAL = 3 + len(SERVICES) + len(GROUPED) + 2


# ================================================================= slides ===
def cover(d):
    r = d.r
    r.slide()
    d.n += 1
    r.rect(0, 0, PW, 7, NAVY)
    r.rect(0, 7, PW, 3, GOLD)
    r.rect(PW - 240, 10, 240, PH - 10, SURFACE)
    r.image(LOGO, MX, 62, 32)
    r.rect(MX, 168, 44, 3, GOLD)
    r.text(MX, 188, 500, 14, [R("FOR PRIVATE HOSPITALS", 10.5, "sansb", TEAL, track=2.6)])
    r.text(MX, 218, 520, 130, [R("What we do for private hospitals", 34, "serifb", INK, ls=1.2)])
    r.text(MX, 330, 500, 60, [R(
        "Recovering what you are owed, bringing in patients, filling posts and fixing what slows the hospital down. "
        "Run by people who have run hospitals.", 15, "serifi", MUTED, ls=1.35)])
    r.line(MX, 413, PW - 240 - 26, 413)
    for dx, label, value in [(0, "FROM", "Consult for Africa"), (170, "WHERE", "Lagos and Abuja"),
                             (310, "WEB", "consultforafrica.com")]:
        r.text(MX + dx, 428, 160, 12, [R(label, 8, "sansb", TEAL, track=1.9)])
        r.text(MX + dx, 446, 160, 14, [R(value, 10.5, "sans", NAVY)])
    r.text(MX, PH - 96, 500, 14, [R("A basis for a conversation, not an offer of terms.", 9, "sans", MUTED)])


def summary(d):
    d.page("On one page")
    y = d.heading("We take the hardest parts of running a hospital off your desk, and we stay until they work.")
    tw = (W - 3 * 12) / 4
    for i, (v, l) in enumerate([
        ("Free", "first review of your unpaid HMO and corporate claims"),
        ("Results", "recovery is paid from what we collect, nothing on money we do not bring in"),
        ("Operators", "former hospital chief executives and clinical leaders on the ground"),
        ("One team", "take one service or several; most hospitals widen the work once it pays"),
    ]):
        d.stat(MX + i * (tw + 12), y + 4, tw, 100, v, l, vsize=22)
    y += 126
    d.kick("What we do", MX, y)
    col = (W - 24) / 2
    left = [
        ("Claims and debt recovery", "Unpaid HMO, corporate and scheme claims, recovered claim by claim."),
        ("Patients and referrals", "Healthcare marketing that is measured in admissions."),
        ("Turnaround", "Cash, cost and revenue control when the numbers turn."),
        ("Organisational audit", "Ten domains, one report, the decisions in order."),
    ]
    right = [
        ("Strategy and growth", "Service lines, pricing, new sites and structuring."),
        ("People and leadership", "Recruitment, locums, assessment, HR and interim leaders."),
        ("Governance and quality", "Clinical governance, accreditation and board set up."),
        ("Systems and specialists", "Hospital systems, dashboards and diaspora consultants."),
    ]
    d.bullets(left, MX, y + 20, col, size=10.5, gap=6)
    d.bullets(right, MX + col + 24, y + 20, col, size=10.5, gap=6)
    d.callout(PH - 98, "Start with a free claims review, or a forty five minute conversation about the one thing that worries you most.")


def why(d):
    d.page("Why us")
    y = d.heading("The difference is who is standing in the hospital when the plan has to work.")
    y = d.lede("Most advisers diagnose and leave. We have run hospitals, so we diagnose and then stay to deliver.", y)
    r = d.r
    half = (W - 20) / 2
    r.rect(MX, y, half, 26, SURFACE)
    r.text(MX + 16, y, half - 30, 26, [R("THE USUAL ADVISER", 9, "sansb", MUTED, track=1.4)], anchor="m")
    r.rect(MX + half + 20, y, half, 26, NAVY)
    r.text(MX + half + 36, y, half - 30, 26, [R("CONSULT FOR AFRICA", 9, "sansb", GOLD, track=1.4)], anchor="m")
    rows = [
        ("Writes the report, then leaves", "Writes less, then stays to deliver it"),
        ("Partners sell, juniors deliver", "Former hospital chief executives and clinicians do the work"),
        ("Frameworks from other markets", "Built for Nigerian payers, costs and staffing"),
        ("Judged on the document", "Judged on cash collected, patients seen and posts filled"),
        ("Leaves you dependent", "Hands over systems your team can run without us"),
    ]
    yy = y + 34
    for old, new in rows:
        r.rect(MX, yy, half, 50, "#F7F9FB")
        r.text(MX + 16, yy, half - 30, 50, [R(old, 12.5, "sans", MUTED)], anchor="m")
        r.rect(MX + half + 20, yy, half, 50, WHITE)
        r.rect(MX + half + 20, yy, 4, 50, GOLD)
        r.text(MX + half + 36, yy, half - 30, 50, [R(new, 12.5, "sansb", NAVY)], anchor="m")
        r.line(MX + half + 20, yy + 50, MX + W, yy + 50)
        yy += 56


def service(d, s):
    d.page(s["kicker"])
    y = d.heading(s["head"])
    y = d.lede(s["lede"], y, w=W - 250)
    lw = W - 250
    cw = (lw - 14) / 2
    ch = 122
    for i, (h, b) in enumerate(s["points"]):
        cx = MX + (i % 2) * (cw + 14)
        cy = y + (i // 2) * (ch + 12)
        d.card(cx, cy, cw, ch, h, b, fill=SURFACE, hsize=12.5, bsize=11)
    px = MX + lw + 22
    pw = W - lw - 22
    d.r.rect(px, y, pw, 2 * ch + 12, NAVY, rounded=True)
    d.r.text(px + 18, y + 18, pw - 36, 12, [R("YOU NEED THIS IF", 8.5, "sansb", GOLD, track=1.6)])
    yy = y + 44
    for sign in s["signs"]:
        d.r.rect(px + 18, yy + 6, 7, 3, GOLD)
        d.r.text(px + 34, yy, pw - 52, 48, [R(sign, 12, "sans", WHITE, ls=1.35)])
        yy += 58
    d.callout(PH - 98, s["terms"])


def grouped(d, g):
    d.page(g["kicker"])
    y = d.heading(g["head"]) + 12
    cw = (W - 18) / 2
    ch = 132
    for i, (h, b) in enumerate(g["cards"]):
        cx = MX + (i % 2) * (cw + 18)
        cy = y + (i // 2) * (ch + 16)
        d.card(cx, cy, cw, ch, h, b, hsize=14, bsize=12)
    d.callout(PH - 98, g["terms"])


def how(d):
    d.page("How we work")
    y = d.heading("Four ways to engage us, and every one is agreed in writing before we start.") + 10
    tw = (W - 3 * 12) / 4
    for i, (h, b) in enumerate([
        ("Paid from results", "Claims and debt recovery. The first review is free and the fee comes out of what we recover."),
        ("Fixed price project", "An audit, a strategy, a feasibility study or a recruitment search, with a defined output and date."),
        ("Monthly retainer", "Marketing, HR or management support that runs month to month against agreed numbers."),
        ("Embedded team", "A turnaround or a fractional leader inside the hospital for a fixed term."),
    ]):
        d.card(MX + i * (tw + 12), y, tw, 190, h, b, hsize=13.5, bsize=11.5)
    y += 222
    d.kick("What you can expect", MX, y)
    d.bullets([
        ("", "A named lead who has done the job you are asking us to fix."),
        ("", "Confidentiality: your data stays with us and your name stays out of our marketing unless you say otherwise."),
        ("", "A handover, so the improvement outlasts the engagement."),
    ], MX, y + 22, W, size=12.5, gap=8)


def start(d):
    d.page("How to start")
    y = d.heading("Three ways to start, none of which commits you to more.") + 10
    tw = (W - 2 * 14) / 3
    for i, (n, h, b) in enumerate([
        ("01", "A free claims review", "Send us your list of unpaid HMO and corporate claims. We tell you what is recoverable and how long it will take."),
        ("02", "A conversation", "Forty five minutes with a senior partner on the one problem that worries you most. No charge."),
        ("03", "An audit", "A fixed price diagnostic of the whole hospital, and a short list of decisions for the board."),
    ]):
        x = MX + i * (tw + 14)
        d.r.rect(x, y, tw, 170, SURFACE, rounded=True)
        d.r.rect(x, y, 5, 170, GOLD)
        d.r.text(x + 20, y + 16, tw - 36, 30, [R(n, 22, "serifb", GOLD)])
        d.r.text(x + 20, y + 50, tw - 36, 90, [
            R(h, 14, "sansb", NAVY, ls=1.25),
            R(b, 11.5, "sans", BODY, sb=5, ls=1.4),
        ])
    y += 196
    d.r.rect(MX, y, W, 112, NAVY, rounded=True)
    d.r.rect(MX, y, 5, 112, GOLD)
    d.r.text(MX + 24, y + 18, 400, 12, [R("SPEAK TO US", 8.5, "sansb", GOLD, track=1.6)])
    d.r.text(MX + 24, y + 40, 360, 60, [
        R("Ejiro Alechenu", 17, "serifb", WHITE, ls=1.2),
        R("Business Development, Consult for Africa", 10.5, "sans", "#C9D6E0", sb=4),
    ])
    d.r.text(MX + 400, y + 40, W - 424, 60, [
        R("hello@consultforafrica.com", 12, "sansb", WHITE, ls=1.4),
        R("consultforafrica.com", 12, "sans", "#C9D6E0", ls=1.4, sb=2),
        R("Lagos and Abuja", 12, "sans", "#C9D6E0", ls=1.4, sb=2),
    ])


def build(r):
    d = Deck(r, TOTAL)
    cover(d)
    summary(d)
    why(d)
    for s in SERVICES:
        service(d, s)
    for g in GROUPED:
        grouped(d, g)
    how(d)
    start(d)
    assert d.n == TOTAL, (d.n, TOTAL)
    r.save()


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    pdf = OUT_DIR / f"{STEM}.pdf"
    pptx = OUT_DIR / f"{STEM}.pptx"
    build(Pdf(pdf))
    build(Pptx(pptx))
    shutil.copyfile(pdf, SHARE)
    print(pdf)
    print(pptx)
    print(SHARE)
