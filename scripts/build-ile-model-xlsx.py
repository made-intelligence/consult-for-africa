"""
ilé five-year financial model as a live Excel workbook.

MECE by construction:
  Revenue  = four business lines, twelve streams, each stream in exactly one line.
  Costs    = direct (by stream) + site operating (residence one, residence two, mortuary)
             + leases + central. EBITDA = revenue less all four.
  Cash     = EBITDA, add back leases charged, less leases paid, capital spend, plus deposits.

Every input is a named, blue cell on Assumptions and mirrors scripts/ile_model.py, so the
workbook and the deck agree. Monthly from November 2026 (month 1) to October 2031 (month 60).

    python3 scripts/build-ile-model-xlsx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ile_model as M  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Ile - Financial Model.xlsx"

N = M.MONTHS
FIRST = 4                                   # column D is month 1
COLS = [get_column_letter(FIRST + i) for i in range(N)]
LASTC = COLS[-1]

ARIAL = "Arial"
F_IN = Font(name=ARIAL, size=10, color="0000FF")
F_CALC = Font(name=ARIAL, size=10, color="000000")
F_LINK = Font(name=ARIAL, size=10, color="008000")
F_B = Font(name=ARIAL, size=10, bold=True)
F_H = Font(name=ARIAL, size=13, bold=True, color="0F7C63")
F_SEC = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_NOTE = Font(name=ARIAL, size=9, italic=True, color="7A847E")
FILL_SEC = PatternFill("solid", fgColor="0F7C63")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_TOT = PatternFill("solid", fgColor="E3EDE7")
THIN = Side(style="thin", color="C9CFC8")
NGN = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x"'
USD = '$#,##0;($#,##0);"-"'

wb = Workbook()
names = {}


def name(nm, ws, cell):
    ref = "'%s'!$%s$%s" % (ws.title, "".join(ch for ch in cell if ch.isalpha()),
                          "".join(ch for ch in cell if ch.isdigit()))
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    names[nm] = ref


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_H
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_NOTE
    ws.sheet_view.showGridLines = False


def section(ws, row, text, width=8):
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).fill = FILL_SEC
    ws.cell(row=row, column=1, value=text).font = F_SEC


# =============================================================== Assumptions ===
A = wb.active
A.title = "Assumptions"
title(A, "Ilé: assumptions",
      "Blue = input you can change. Every formula in the workbook reads these named cells. "
      "Yellow = scenario lever.")
A.column_dimensions["A"].width = 46
for c in "BCDEFG":
    A.column_dimensions[c].width = 14
A.column_dimensions["H"].width = 70

r = 4


def scalar(label, nm, value, fmt=NGN, unit="", note="", key=False):
    global r
    A.cell(row=r, column=1, value=label).font = F_CALC
    cell = A.cell(row=r, column=2, value=value)
    cell.font = F_IN
    cell.number_format = fmt
    if key:
        cell.fill = FILL_KEY
    A.cell(row=r, column=3, value=unit).font = F_NOTE
    A.cell(row=r, column=8, value=note).font = F_NOTE
    name(nm, A, "B%d" % r)
    r += 1


def formula_cell(label, nm, f, fmt=NGN, note=""):
    global r
    A.cell(row=r, column=1, value=label).font = F_CALC
    cell = A.cell(row=r, column=2, value=f)
    cell.font = F_CALC
    cell.number_format = fmt
    A.cell(row=r, column=8, value=note).font = F_NOTE
    name(nm, A, "B%d" % r)
    r += 1


def points(label, prefix, pts, fmt=NGN, note=""):
    """Ramp or step points: month and value pairs in B..G, named PREFIX_M1, PREFIX_V1 ..."""
    global r
    A.cell(row=r, column=1, value=label).font = F_CALC
    for k, (m, v) in enumerate(pts):
        cm = A.cell(row=r, column=2 + 2 * k, value=m)
        cm.font = F_IN
        cm.number_format = "0"
        cv = A.cell(row=r, column=3 + 2 * k, value=v)
        cv.font = F_IN
        cv.number_format = fmt
        name("%s_M%d" % (prefix, k + 1), A, "%s%d" % (get_column_letter(2 + 2 * k), r))
        name("%s_V%d" % (prefix, k + 1), A, "%s%d" % (get_column_letter(3 + 2 * k), r))
    A.cell(row=r, column=8, value=note).font = F_NOTE
    r += 1


def head_points():
    global r
    for k, t in enumerate(["Month 1", "Value 1", "Month 2", "Value 2", "Month 3", "Value 3"]):
        A.cell(row=r, column=2 + k, value=t).font = F_NOTE
    r += 1


section(A, r, "0. Model and scenario levers"); r += 1
scalar("Naira per US dollar", "FX", M.FX, NGN, "₦/$", "Flat over five years; families abroad pay in dollars or pounds")
scalar("Days in a month", "DAYS_PM", M.DAYS, NUM1, "days")
scalar("Residence fill speed", "RES_SPEED", 1, PCT, "x base", "1 = base case; 0.5 = half speed (applies to all three wings)", key=True)
scalar("Mortuary fill speed", "MORT_SPEED", 1, PCT, "x base", "1 = base case; 0.5 = half speed", key=True)
scalar("Residence delay", "RES_DELAY", 0, "0", "months", "Months both residences open late", key=True)
scalar("Software licences sold", "SAAS_ON", 1, "0", "1 or 0", "0 = no software revenue", key=True)
r += 1

section(A, r, "1. Care in the home"); r += 1
head_points()
points("Care at home clients (ramp)", "HOME", M.HOME_CLIENTS, "#,##0", "Clients on subscription; linear between points, flat after")
scalar("Care at home revenue per client", "HOME_ARPU", M.HOME_ARPU, NGN, "₦/month", "Blended subscription; Lagos agencies charge ₦119k to ₦359k for carer and nurse packages")
scalar("Care at home direct cost", "HOME_DIRECT", M.HOME_DIRECT, PCT, "of revenue", "Carer pay, travel, consumables")
points("Concierge members (ramp)", "CONC", M.CONCIERGE, "#,##0")
scalar("Concierge fee", "CONC_FEE", M.CONCIERGE_FEE, NGN, "₦/month", "Named doctor and care manager, 24-hour line, hospital visits")
scalar("Concierge direct cost", "CONC_DIRECT", M.CONCIERGE_DIRECT, PCT, "of revenue")
points("Virtual consultations a month (ramp)", "CONS", M.CONSULTS, "#,##0")
scalar("Consultation fee", "CONS_FEE", M.CONSULT_FEE, NGN, "₦ each")
scalar("Consultation direct cost", "CONS_DIRECT", M.CONSULT_DIRECT, PCT, "of revenue")
points("Post-operative rehab packages a month (ramp)", "REHAB", M.REHAB, "#,##0")
scalar("Rehab package fee", "REHAB_FEE", M.REHAB_FEE, NGN, "₦ per six weeks")
scalar("Rehab direct cost", "REHAB_DIRECT", M.REHAB_DIRECT, PCT, "of revenue")
points("Palliative patients at home (ramp)", "PALLH", M.PALL_HOME, "#,##0")
scalar("Palliative at home fee", "PALLH_FEE", M.PALL_HOME_ARPU, NGN, "₦/month")
scalar("Palliative at home direct cost", "PALLH_DIRECT", M.PALL_HOME_DIRECT, PCT, "of revenue")
r += 1

section(A, r, "2. Residential care (per residence)"); r += 1
scalar("Residence one opens", "RES1_OPEN", M.RESIDENCES[0], "0", "month", "Month 15 = January 2028")
scalar("Residence two opens", "RES2_OPEN", M.RESIDENCES[1], "0", "month", "Month 43 = May 2030; paid for from cash flow")
formula_cell("Residence one opens, after any delay", "RES1_EFF", "=RES1_OPEN+RES_DELAY", "0")
formula_cell("Residence two opens, after any delay", "RES2_EFF", "=RES2_OPEN+RES_DELAY", "0")
scalar("Long-stay beds", "RES_BEDS", M.RES_BEDS, "0", "beds", "20 beds in all: 12 long-stay, 4 step-down, 4 palliative")
scalar("Step-down and rehab beds", "STEP_BEDS", M.STEP_BEDS, "0", "beds")
scalar("Palliative beds", "PALL_BEDS", M.PALL_BEDS, "0", "beds")
scalar("Long-stay occupied at opening", "RES_OPEN_OCC", M.RES_OPEN_OCC, NUM1, "beds", "Reserved by deposit before opening")
scalar("Long-stay beds added a month", "RES_STEP", M.RES_STEP, NUM1, "beds")
scalar("Long-stay occupancy ceiling", "RES_CAP", M.RES_CAP, NUM1, "beds", "11 of 12, 92%")
scalar("Step-down occupied at opening", "STEP_OPEN_OCC", M.STEP_OPEN_OCC, NUM1, "beds")
scalar("Step-down beds added a month", "STEP_STEP", M.STEP_STEP, NUM1, "beds")
scalar("Step-down occupancy ceiling", "STEP_CAP", M.STEP_CAP, NUM1, "beds")
scalar("Palliative occupied at opening", "PALL_OPEN_OCC", M.PALL_OPEN_OCC, NUM1, "beds")
scalar("Palliative beds added a month", "PALL_STEP", M.PALL_STEP, NUM1, "beds")
scalar("Palliative occupancy ceiling", "PALL_CAP", M.PALL_CAP, NUM1, "beds")
scalar("Long-stay fee", "RES_FEE", M.RES_FEE, NGN, "₦/month", "About $1,070; Lagos homes publish ₦79k to ₦249k without night nursing")
scalar("Step-down day rate", "STEP_DAY", M.STEP_DAY, NGN, "₦/day", "Nursing and daily physiotherapy")
scalar("Step-down average stay", "STEP_STAY", M.STEP_STAY, "0", "days")
scalar("Palliative wing fee", "PALL_FEE", M.PALL_FEE, NGN, "₦/month")
scalar("Staffing, loaded", "RES_STAFF", M.RES_STAFF, NGN, "₦/month", "Nurses on every shift, care assistants, manager, therapy; includes on-costs")
scalar("Pre-opening staffing", "PREOPEN", 0.6, PCT, "of staffing", "Charged in each of the two months before opening")
scalar("Food, consumables, laundry per occupied bed", "RES_PER_BED", M.RES_PER_BED, NGN, "₦/month")
scalar("Power and diesel", "RES_UTIL", M.RES_UTIL, NGN, "₦/month")
scalar("Insurance, maintenance, licences", "RES_OTHER", M.RES_OTHER, NGN, "₦/month")
scalar("Building lease", "RES_LEASE_YR", M.RES_RENT * 12, NGN, "₦/year", "Unconverted building in Lekki; paid a year ahead from signing; estimate")
scalar("Lease signed before opening", "LEASE_LEAD", M.LEASE_LEAD, "0", "months", "Building works run in these months")
scalar("Building works", "RES_CONV", M.RES_CONVERSION, NGN, "₦", "Conversion to a licensable care home; spread over the lease lead")
scalar("Equipment, furnishing, nurse call, power, pre-opening", "RES_EQUIP", M.RES_CAPEX, NGN, "₦")
scalar("Equipment spend spread over", "EQUIP_MONTHS", 5, "0", "months before opening")
scalar("Deposit per reserved place", "DEPOSIT", M.DEPOSIT, NGN, "₦", "One month's fee, held not earned")
scalar("Places reserved by deposit at opening", "DEPOSITS_N", M.DEPOSITS_AT_OPEN, "0", "places")
r += 1

section(A, r, "3. Mortuary and pathology"); r += 1
scalar("Mortuary opens", "MORT_OPEN", M.MORT_OPEN, "0", "month", "Month 10 = August 2027; Lekki corridor")
scalar("Places at opening", "MORT_CAP1", M.MORT_CAP[0][1], "0", "places")
scalar("Expansion month", "MORT_EXP_M", M.MORT_CAP[1][0], "0", "month", "Month 34 = August 2029")
scalar("Places after expansion", "MORT_CAP2", M.MORT_CAP[1][1], "0", "places")
scalar("Occupancy ceiling", "MORT_FILL", M.MORT_FILL, PCT, "of places")
scalar("Bodies held at opening", "MORT_HELD_OPEN", M.MORT_HELD_OPEN, NUM1, "bodies")
scalar("Bodies held added a month", "MORT_HELD_STEP", M.MORT_HELD_STEP, NUM1, "bodies")
scalar("Average stay", "MORT_STAY", M.MORT_STAY, "0", "days", "Lagos families hold four to eight weeks")
scalar("Preservation day rate", "MORT_DAY", M.MORT_DAY, NGN, "₦/day", "Private homes charge ₦15,000 to ₦50,000 a day")
scalar("Preparation fee per family", "MORT_CASE", M.MORT_CASE, NGN, "₦", "Embalming, dressing, documentation, local transport")
scalar("Consumables per family", "MORT_CASE_COST", M.MORT_CASE_COST, NGN, "₦")
head_points()
points("Remains received from abroad a month (ramp)", "REPAT", M.REPAT, "#,##0")
scalar("Repatriation fee", "REPAT_FEE", M.REPAT_FEE, NGN, "₦")
scalar("Repatriation direct cost", "REPAT_DIRECT", M.REPAT_DIRECT, PCT, "of revenue")
points("Autopsies a month (ramp)", "AUT", M.AUTOPSY, "#,##0")
scalar("Autopsy fee", "AUT_FEE", M.AUTOPSY_FEE, NGN, "₦", "Estimate; no published Lagos price found")
scalar("Autopsy direct cost", "AUT_DIRECT", M.AUTOPSY_DIRECT, PCT, "of revenue", "Sessional pathologist, histology, toxicology")
head_points()
points("Mortuary staffing (step)", "MSTAFF", M.MORT_STAFF, NGN, "₦/month from each month")
points("Mortuary power (step)", "MPOWER", M.MORT_POWER, NGN, "₦/month; cold storage runs 24 hours")
scalar("Mortuary other running costs", "MORT_OTHER", M.MORT_OTHER, NGN, "₦/month")
scalar("Mortuary site lease", "MORT_LEASE_YR", M.MORT_RENT * 12, NGN, "₦/year", "Paid a year ahead; estimate")
scalar("Lease signed", "MORT_LEASE_FROM", M.MORT_LEASE_FROM, "0", "month")
scalar("Build starts", "MBUILD_M", M.MORT_CAPEX[0][0], "0", "month")
scalar("Build: building, cold storage, vehicles", "MBUILD", sum(v for m, v in M.MORT_CAPEX[:3]), NGN, "₦")
scalar("Build spread over", "MBUILD_N", 3, "0", "months")
scalar("Expansion build starts", "MEXP_M", M.MORT_CAPEX[3][0], "0", "month")
scalar("Expansion build", "MEXP", sum(v for m, v in M.MORT_CAPEX[3:]), NGN, "₦")
scalar("Expansion spread over", "MEXP_N", 3, "0", "months")
scalar("Autopsy suite, month", "AUTCAP_M", M.AUTOPSY_CAPEX[0], "0", "month")
scalar("Autopsy suite", "AUTCAP", M.AUTOPSY_CAPEX[1], NGN, "₦")
r += 1

section(A, r, "4. Software"); r += 1
head_points()
points("Organisations licensing the platform (ramp)", "SAAS", M.SAAS, "#,##0", "Agencies, care homes and hospices")
scalar("Licence fee", "SAAS_FEE", M.SAAS_FEE, NGN, "₦/month per organisation")
scalar("Hosting, onboarding and support", "SAAS_DIRECT", M.SAAS_DIRECT, PCT, "of revenue")
r += 1

section(A, r, "5. Central overhead"); r += 1
head_points()
points("Central team (step)", "CENT", M.CENTRAL, NGN, "₦/month from each month")
points("Marketing (step)", "MKT", M.MARKETING, NGN, "₦/month from each month")
r += 1

section(A, r, "6. The round and exit"); r += 1
scalar("Raise", "RAISE_USD", 500_000, USD, "$", "Pre-seed")
scalar("Stake offered", "STAKE", 0.30, PCT, "", "", key=True)
formula_cell("Raise in naira", "RAISE_NGN", "=RAISE_USD*FX")
formula_cell("Post-money", "POST_USD", "=RAISE_USD/STAKE", USD)
formula_cell("Pre-money", "PRE_USD", "=POST_USD-RAISE_USD", USD)
scalar("Exit multiple, low", "EXIT_LO", 6, MULT, "x EBITDA", "Discount to SCI 12.1x and Amedisys ~13x")
scalar("Exit multiple, mid", "EXIT_MID", 8, MULT, "x EBITDA")
scalar("Exit multiple, high", "EXIT_HI", 10, MULT, "x EBITDA")
scalar("Naira at exit, weak case", "FX_WEAK", 2000, NGN, "₦/$")
r += 1

section(A, r, "7. Market sizing, Lagos"); r += 1
scalar("Lagos population", "LAGOS_POP", M.LAGOS_POP, NGN, "people", "Estimates run 12.8m to 17.2m")
scalar("Share aged 60 and over", "SHARE_60", M.SHARE_60, PCT, "", "Nigeria is about 7%; Lagos is younger")
scalar("Dependent in daily living", "ADL_DEP", M.ADL_DEP, PCT, "", "Nigerian study of people aged 60+")
scalar("Of the dependent, needing 24-hour care", "NEEDS_24H", M.NEEDS_24H, PCT)
scalar("Crude death rate", "DEATH_RATE", M.DEATH_RATE, '0.000', "per person a year")
scalar("Deaths that need palliative care", "PALL_NEED", M.PALL_NEED, PCT)
scalar("Palliative months per death", "PALL_MONTHS", M.PALL_MONTHS, "0", "months")
scalar("Deaths kept in a private funeral home", "PRIVATE_MORT", M.PRIVATE_MORT, PCT)
scalar("Families able to pay these prices", "PREMIUM", M.PREMIUM, PCT, "", "Or funded from abroad")
A.freeze_panes = "B4"


# ============================================================ monthly sheets ===
def monthly_sheet(title_text, sub):
    ws = wb.create_sheet(title_text)
    title(ws, "Ilé: " + title_text.lower(), sub)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 3
    for c in COLS:
        ws.column_dimensions[c].width = 11.5
    ws.cell(row=4, column=1, value="Month").font = F_B
    ws.cell(row=5, column=1, value="Period").font = F_B
    ws.cell(row=6, column=1, value="Model year").font = F_B
    for i, c in enumerate(COLS):
        ws["%s4" % c] = i + 1
        ws["%s4" % c].font = F_B
        ws["%s5" % c] = M.label(i + 1)
        ws["%s5" % c].font = F_B
        ws["%s5" % c].alignment = Alignment(horizontal="right")
        ws["%s6" % c] = "=INT((%s$4-1)/12)+1" % c
        ws["%s6" % c].font = F_CALC
    ws.freeze_panes = "D7"
    return ws


def ramp(prefix, m):
    p = prefix
    return ("IF({m}<{p}_M1,0,IF({m}<={p}_M2,{p}_V1+({p}_V2-{p}_V1)*({m}-{p}_M1)/({p}_M2-{p}_M1),"
            "IF({m}<={p}_M3,{p}_V2+({p}_V3-{p}_V2)*({m}-{p}_M2)/({p}_M3-{p}_M2),{p}_V3)))").format(p=p, m=m)


def step2(prefix, m):
    return "IF({m}>={p}_M2,{p}_V2,IF({m}>={p}_M1,{p}_V1,0))".format(p=prefix, m=m)


def step3(prefix, m):
    return "IF({m}>={p}_M3,{p}_V3,IF({m}>={p}_M2,{p}_V2,IF({m}>={p}_M1,{p}_V1,0)))".format(p=prefix, m=m)


rows = {}          # (sheet, key) -> row number


def line(ws, row, key, label, tmpl, fmt=NGN, bold=False, total=False, unit=""):
    """tmpl uses {c} for this column and {m} for this month's index cell."""
    ws.cell(row=row, column=1, value=label).font = F_B if bold else F_CALC
    ws.cell(row=row, column=2, value=unit).font = F_NOTE
    for c in COLS:
        cell = ws["%s%d" % (c, row)]
        cell.value = "=" + tmpl.format(c=c, m="%s$4" % c)
        cell.font = F_B if bold else F_CALC
        cell.number_format = fmt
        if total:
            cell.fill = FILL_TOT
            cell.border = Border(top=THIN)
    if total:
        ws.cell(row=row, column=1).fill = FILL_TOT
    rows[(ws.title, key)] = row


def ref(sheet, key, c):
    return "'%s'!%s%d" % (sheet, c, rows[(sheet, key)])


def heading_row(ws, row, text):
    section(ws, row, text, 3)


# ------------------------------------------------------------------ Volumes ---
V = monthly_sheet("Volumes", "Every count that drives revenue and cost. Linear ramps between the "
                             "points on Assumptions.")
row = 8
heading_row(V, row, "1. Care in the home"); row += 1
line(V, row, "home", "Care at home clients", ramp("HOME", "{m}"), NUM1, unit="clients"); row += 1
line(V, row, "conc", "Concierge members", ramp("CONC", "{m}"), NUM1, unit="members"); row += 1
line(V, row, "cons", "Virtual consultations", ramp("CONS", "{m}"), NUM1, unit="a month"); row += 1
line(V, row, "rehab", "Post-operative rehab packages begun", ramp("REHAB", "{m}"), NUM1, unit="a month"); row += 1
line(V, row, "pallh", "Palliative patients at home", ramp("PALLH", "{m}"), NUM1, unit="patients"); row += 2
heading_row(V, row, "2. Residential care, occupied beds"); row += 1
for rn, op in [("1", "RES1_EFF"), ("2", "RES2_EFF")]:
    for key, lab, o, s, cap in [("ls", "long-stay", "RES_OPEN_OCC", "RES_STEP", "RES_CAP"),
                                ("sd", "step-down", "STEP_OPEN_OCC", "STEP_STEP", "STEP_CAP"),
                                ("pw", "palliative", "PALL_OPEN_OCC", "PALL_STEP", "PALL_CAP")]:
        line(V, row, "r%s%s" % (rn, key), "Residence %s, %s" % ("one" if rn == "1" else "two", lab),
             "IF({m}>=%s,MIN(%s,%s+%s*RES_SPEED*({m}-%s)),0)" % (op, cap, o, s, op), NUM1, unit="beds")
        row += 1
line(V, row, "beds", "Occupied beds, both residences",
     "SUM({c}%d:{c}%d)" % (row - 6, row - 1), NUM1, bold=True, total=True, unit="beds"); row += 2
heading_row(V, row, "3. Mortuary and pathology"); row += 1
line(V, row, "mcap", "Places", "IF({m}>=MORT_EXP_M,MORT_CAP2,IF({m}>=MORT_OPEN,MORT_CAP1,0))", "#,##0", unit="places"); row += 1
line(V, row, "held", "Bodies held on an average day",
     "IF({m}>=MORT_OPEN,MIN({c}%d*MORT_FILL,MORT_HELD_OPEN+MORT_HELD_STEP*MORT_SPEED*({m}-MORT_OPEN)),0)" % (row - 1),
     NUM1, unit="bodies"); row += 1
line(V, row, "cases", "Families served", "{c}%d*DAYS_PM/MORT_STAY" % (row - 1), NUM1, unit="a month"); row += 1
line(V, row, "repat", "Remains received from abroad", "IF({m}>=MORT_OPEN,%s,0)" % ramp("REPAT", "{m}"), NUM1, unit="a month"); row += 1
line(V, row, "aut", "Autopsies", "IF({m}>=MORT_OPEN,%s,0)" % ramp("AUT", "{m}"), NUM1, unit="a month"); row += 2
heading_row(V, row, "4. Software"); row += 1
line(V, row, "orgs", "Organisations licensing", "%s*SAAS_ON" % ramp("SAAS", "{m}"), NUM1, unit="orgs"); row += 1


def v(key):
    return "'Volumes'!{c}%d" % rows[("Volumes", key)]


# ------------------------------------------------------------------ Revenue ---
R = monthly_sheet("Revenue", "Twelve streams in four business lines. Naira a month.")
row = 8
heading_row(R, row, "1. Care in the home"); row += 1
s1 = row
line(R, row, "conc", "Concierge geriatric care", v("conc") + "*CONC_FEE"); row += 1
line(R, row, "home", "Care at home", v("home") + "*HOME_ARPU"); row += 1
line(R, row, "cons", "Virtual consultations", v("cons") + "*CONS_FEE"); row += 1
line(R, row, "rehab", "Post-operative rehabilitation", v("rehab") + "*REHAB_FEE"); row += 1
line(R, row, "pallh", "Palliative care at home", v("pallh") + "*PALLH_FEE"); row += 1
line(R, row, "L1", "Care in the home", "SUM({c}%d:{c}%d)" % (s1, row - 1), bold=True, total=True); row += 2
heading_row(R, row, "2. Residential care"); row += 1
s2 = row
line(R, row, "long", "Residence, long-stay", "(%s+%s)*RES_FEE" % (v("r1ls"), v("r2ls"))); row += 1
line(R, row, "step", "Step-down beds", "(%s+%s)*DAYS_PM*STEP_DAY" % (v("r1sd"), v("r2sd"))); row += 1
line(R, row, "pall", "Palliative wing", "(%s+%s)*PALL_FEE" % (v("r1pw"), v("r2pw"))); row += 1
line(R, row, "L2", "Residential care", "SUM({c}%d:{c}%d)" % (s2, row - 1), bold=True, total=True); row += 2
heading_row(R, row, "3. Mortuary and pathology"); row += 1
s3 = row
line(R, row, "pres", "Mortuary preservation", v("held") + "*DAYS_PM*MORT_DAY"); row += 1
line(R, row, "prep", "Preparation and repatriation",
     "%s*MORT_CASE+%s*REPAT_FEE" % (v("cases"), v("repat"))); row += 1
line(R, row, "aut", "Autopsy and pathology", v("aut") + "*AUT_FEE"); row += 1
line(R, row, "L3", "Mortuary and pathology", "SUM({c}%d:{c}%d)" % (s3, row - 1), bold=True, total=True); row += 2
heading_row(R, row, "4. Software"); row += 1
line(R, row, "saas", "Software licences", v("orgs") + "*SAAS_FEE"); row += 1
line(R, row, "L4", "Software", "{c}%d" % (row - 1), bold=True, total=True); row += 2
line(R, row, "total", "Revenue", "{c}%d+{c}%d+{c}%d+{c}%d" % tuple(rows[("Revenue", k)] for k in ("L1", "L2", "L3", "L4")),
     bold=True, total=True)


def rv(key):
    return "'Revenue'!{c}%d" % rows[("Revenue", key)]


# -------------------------------------------------------------------- Costs ---
C = monthly_sheet("Costs", "Four buckets: direct, site operating, leases, central. Naira a month.")
row = 8
heading_row(C, row, "A. Direct costs, by stream"); row += 1
sA = row
for key, lab, f in [("conc", "Concierge", rv("conc") + "*CONC_DIRECT"),
                    ("home", "Care at home", rv("home") + "*HOME_DIRECT"),
                    ("cons", "Virtual consultations", rv("cons") + "*CONS_DIRECT"),
                    ("rehab", "Post-operative rehabilitation", rv("rehab") + "*REHAB_DIRECT"),
                    ("pallh", "Palliative care at home", rv("pallh") + "*PALLH_DIRECT"),
                    ("prep", "Mortuary consumables per family", v("cases") + "*MORT_CASE_COST"),
                    ("repat", "Repatriation", v("repat") + "*REPAT_FEE*REPAT_DIRECT"),
                    ("aut", "Autopsy and pathology", rv("aut") + "*AUT_DIRECT"),
                    ("saas", "Software hosting and support", rv("saas") + "*SAAS_DIRECT")]:
    line(C, row, "d_" + key, lab, f); row += 1
line(C, row, "DIRECT", "Direct costs", "SUM({c}%d:{c}%d)" % (sA, row - 1), bold=True, total=True); row += 2
heading_row(C, row, "B. Site operating costs"); row += 1
sB = row
for rn, op in [("1", "RES1_EFF"), ("2", "RES2_EFF")]:
    nm = "one" if rn == "1" else "two"
    line(C, row, "r%s_staff" % rn, "Residence %s, staffing" % nm,
         "IF({m}>=%s,RES_STAFF,IF(OR({m}=%s-1,{m}=%s-2),RES_STAFF*PREOPEN,0))" % (op, op, op)); row += 1
    line(C, row, "r%s_bed" % rn, "Residence %s, food and consumables" % nm,
         "(%s+%s+%s)*RES_PER_BED" % (v("r%sls" % rn), v("r%ssd" % rn), v("r%spw" % rn))); row += 1
    line(C, row, "r%s_util" % rn, "Residence %s, power and diesel" % nm, "IF({m}>=%s,RES_UTIL,0)" % op); row += 1
    line(C, row, "r%s_oth" % rn, "Residence %s, insurance and maintenance" % nm, "IF({m}>=%s,RES_OTHER,0)" % op); row += 1
line(C, row, "m_staff", "Mortuary, staffing", step2("MSTAFF", "{m}")); row += 1
line(C, row, "m_power", "Mortuary, power", step2("MPOWER", "{m}")); row += 1
line(C, row, "m_oth", "Mortuary, other running costs", "IF({m}>=MORT_OPEN,MORT_OTHER,0)"); row += 1
line(C, row, "SITE", "Site operating costs", "SUM({c}%d:{c}%d)" % (sB, row - 1), bold=True, total=True); row += 2
heading_row(C, row, "C. Leases, charged monthly"); row += 1
sC = row
line(C, row, "l_r1", "Residence one building", "IF({m}>=RES1_EFF-LEASE_LEAD,RES_LEASE_YR/12,0)"); row += 1
line(C, row, "l_r2", "Residence two building", "IF({m}>=RES2_EFF-LEASE_LEAD,RES_LEASE_YR/12,0)"); row += 1
line(C, row, "l_m", "Mortuary site", "IF({m}>=MORT_LEASE_FROM,MORT_LEASE_YR/12,0)"); row += 1
line(C, row, "LEASE", "Leases", "SUM({c}%d:{c}%d)" % (sC, row - 1), bold=True, total=True); row += 2
heading_row(C, row, "D. Central overhead"); row += 1
sD = row
line(C, row, "cent", "Central team", step3("CENT", "{m}")); row += 1
line(C, row, "mkt", "Marketing", step2("MKT", "{m}")); row += 1
line(C, row, "CENTRAL", "Central overhead", "SUM({c}%d:{c}%d)" % (sD, row - 1), bold=True, total=True); row += 2
line(C, row, "TOTAL", "Operating costs", "{c}%d+{c}%d+{c}%d+{c}%d" % tuple(rows[("Costs", k)] for k in ("DIRECT", "SITE", "LEASE", "CENTRAL")),
     bold=True, total=True)


def cs(key):
    return "'Costs'!{c}%d" % rows[("Costs", key)]


# ---------------------------------------------------------------------- P&L ---
P = monthly_sheet("P&L", "Revenue by line, the four cost buckets, EBITDA. Before tax and depreciation.")
row = 8
heading_row(P, row, "Revenue"); row += 1
for k, lab in [("L1", "Care in the home"), ("L2", "Residential care"), ("L3", "Mortuary and pathology"),
               ("L4", "Software")]:
    line(P, row, k, lab, rv(k)); P.cell(row=row, column=1).font = F_LINK; row += 1
line(P, row, "rev", "Revenue", "SUM({c}%d:{c}%d)" % (row - 4, row - 1), bold=True, total=True); row += 2
heading_row(P, row, "Costs"); row += 1
for k, lab in [("DIRECT", "Direct costs"), ("SITE", "Site operating costs"), ("LEASE", "Leases"),
               ("CENTRAL", "Central overhead")]:
    line(P, row, k, lab, cs(k)); row += 1
line(P, row, "cost", "Operating costs", "SUM({c}%d:{c}%d)" % (row - 4, row - 1), bold=True, total=True); row += 2
line(P, row, "contrib", "Contribution after direct costs", "{c}%d-{c}%d" % (rows[("P&L", "rev")], rows[("P&L", "DIRECT")])); row += 1
line(P, row, "ebitda", "EBITDA", "{c}%d-{c}%d" % (rows[("P&L", "rev")], rows[("P&L", "cost")]), bold=True, total=True); row += 1
line(P, row, "margin", "EBITDA margin", "IF({c}%d=0,0,{c}%d/{c}%d)" % (rows[("P&L", "rev")], rows[("P&L", "ebitda")], rows[("P&L", "rev")]), PCT); row += 1
line(P, row, "posflag", "Positive for six months running (flag)",
     "IF(COUNTIF({c}%d:%s%d,\">0\")=MIN(6,%d-{m}+1),IF({c}%d>0,1,0),0)" % (
         rows[("P&L", "ebitda")], "{{c6}}", rows[("P&L", "ebitda")], N, rows[("P&L", "ebitda")]), "0")
# the six-month window end column, per month
for i, c in enumerate(COLS):
    c6 = COLS[min(i + 5, N - 1)]
    cell = P["%s%d" % (c, rows[("P&L", "posflag")])]
    cell.value = cell.value.replace("{c6}", c6)


def pl(key):
    return "'P&L'!{c}%d" % rows[("P&L", key)]


# ---------------------------------------------------------------- Cash Flow ---
F = monthly_sheet("Cash Flow", "EBITDA to cash. Before tax. The raise arrives in month 1.")
row = 8
line(F, row, "ebitda", "EBITDA", pl("ebitda")); F.cell(row=row, column=1).font = F_LINK; row += 1
line(F, row, "addl", "Add back leases charged", cs("LEASE")); row += 1
heading_row(F, row, "Leases paid a year ahead"); row += 1
sL = row
line(F, row, "p_r1", "Residence one building",
     "-IF(AND({m}>=RES1_EFF-LEASE_LEAD,MOD({m}-(RES1_EFF-LEASE_LEAD),12)=0),RES_LEASE_YR,0)"); row += 1
line(F, row, "p_r2", "Residence two building",
     "-IF(AND({m}>=RES2_EFF-LEASE_LEAD,MOD({m}-(RES2_EFF-LEASE_LEAD),12)=0),RES_LEASE_YR,0)"); row += 1
line(F, row, "p_m", "Mortuary site",
     "-IF(AND({m}>=MORT_LEASE_FROM,MOD({m}-MORT_LEASE_FROM,12)=0),MORT_LEASE_YR,0)"); row += 1
line(F, row, "LPAID", "Leases paid", "SUM({c}%d:{c}%d)" % (sL, row - 1), bold=True, total=True); row += 2
heading_row(F, row, "Capital spend"); row += 1
sK = row
for rn, op in [("1", "RES1_EFF"), ("2", "RES2_EFF")]:
    nm = "one" if rn == "1" else "two"
    line(F, row, "k_r%s_w" % rn, "Residence %s, building works" % nm,
         "-IF(AND({m}>=%s-LEASE_LEAD,{m}<%s),RES_CONV/LEASE_LEAD,0)" % (op, op)); row += 1
    line(F, row, "k_r%s_e" % rn, "Residence %s, equipment and pre-opening" % nm,
         "-IF(AND({m}>=%s-EQUIP_MONTHS,{m}<%s),RES_EQUIP/EQUIP_MONTHS,0)" % (op, op)); row += 1
line(F, row, "k_mb", "Mortuary build", "-IF(AND({m}>=MBUILD_M,{m}<MBUILD_M+MBUILD_N),MBUILD/MBUILD_N,0)"); row += 1
line(F, row, "k_aut", "Autopsy suite", "-IF({m}=AUTCAP_M,AUTCAP,0)"); row += 1
line(F, row, "k_me", "Mortuary expansion", "-IF(AND({m}>=MEXP_M,{m}<MEXP_M+MEXP_N),MEXP/MEXP_N,0)"); row += 1
line(F, row, "CAPEX", "Capital spend", "SUM({c}%d:{c}%d)" % (sK, row - 1), bold=True, total=True); row += 2
line(F, row, "dep", "Residence deposits received",
     "IF({m}=RES1_EFF-1,DEPOSIT*DEPOSITS_N,0)+IF({m}=RES2_EFF-1,DEPOSIT*DEPOSITS_N,0)"); row += 2
line(F, row, "net", "Net cash flow before funding",
     "{c}%d+{c}%d+{c}%d+{c}%d+{c}%d" % tuple(rows[("Cash Flow", k)] for k in ("ebitda", "addl", "LPAID", "CAPEX", "dep")),
     bold=True, total=True); row += 1
r_net = rows[("Cash Flow", "net")]
line(F, row, "cum", "Cumulative cash before funding", "SUM($D$%d:{c}%d)" % (r_net, r_net), bold=True); row += 1
line(F, row, "raise", "Equity raised", "IF({m}=1,RAISE_NGN,0)"); row += 1
line(F, row, "bal", "Cash balance", "{c}%d+SUM($D$%d:{c}%d)" % (rows[("Cash Flow", "cum")], rows[("Cash Flow", "raise")],
                                                              rows[("Cash Flow", "raise")]), bold=True, total=True); row += 2

r_cum = rows[("Cash Flow", "cum")]
kpi = row
F.cell(row=kpi, column=1, value="Peak funding need").font = F_B
F.cell(row=kpi, column=2, value="=-MIN($D$%d:$%s$%d)" % (r_cum, LASTC, r_cum)).number_format = NGN
name("PEAK", F, "B%d" % kpi)
F.cell(row=kpi + 1, column=1, value="Month cash is lowest").font = F_B
F.cell(row=kpi + 1, column=2, value="=INDEX($D$5:$%s$5,MATCH(MIN($D$%d:$%s$%d),$D$%d:$%s$%d,0))" % (LASTC, r_cum, LASTC, r_cum, r_cum, LASTC, r_cum))
F.cell(row=kpi + 2, column=1, value="EBITDA positive, and stays positive, from").font = F_B
r_flag = rows[("P&L", "posflag")]
F.cell(row=kpi + 2, column=2, value="=IFERROR(INDEX($D$5:$%s$5,MATCH(1,'P&L'!$D$%d:$%s$%d,0)),\"Not in five years\")" % (LASTC, r_flag, LASTC, r_flag))
F.cell(row=kpi + 3, column=1, value="Lowest cash balance after the raise").font = F_B
r_bal = rows[("Cash Flow", "bal")]
F.cell(row=kpi + 3, column=2, value="=MIN($D$%d:$%s$%d)" % (r_bal, LASTC, r_bal)).number_format = NGN
name("MINBAL", F, "B%d" % (kpi + 3))
for k in range(4):
    F.cell(row=kpi + k, column=2).font = F_B
name("TROUGH_M", F, "B%d" % (kpi + 1))
name("EBITDA_POS", F, "B%d" % (kpi + 2))


# ---------------------------------------------------------------- Annual -------
Y = wb.create_sheet("Annual")
title(Y, "Ilé: five-year summary", "Year one is November 2026 to October 2027. Naira, before tax. "
                                    "Every figure sums the monthly sheets.")
Y.column_dimensions["A"].width = 44
for c in "BCDEFG":
    Y.column_dimensions[c].width = 15
hdr = ["", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
for k, t in enumerate(hdr):
    cell = Y.cell(row=4, column=1 + k, value=t)
    cell.font = F_SEC
    cell.fill = FILL_SEC
    cell.alignment = Alignment(horizontal="right" if k else "left")
Y.cell(row=5, column=1, value="Period").font = F_NOTE
for y in range(5):
    Y.cell(row=5, column=2 + y, value="%s to %s" % (M.label(y * 12 + 1), M.label(y * 12 + 12))).font = F_NOTE
yrow = [6]
ann = {}


def annual(label, sheet, key, bold=False, fmt=NGN):
    rr = yrow[0]
    Y.cell(row=rr, column=1, value=label).font = F_B if bold else F_CALC
    src = rows[(sheet, key)]
    for y in range(5):
        col = get_column_letter(2 + y)
        Y["%s%d" % (col, rr)] = "=SUMIF('%s'!$D$6:$%s$6,%d,'%s'!$D$%d:$%s$%d)" % (sheet, LASTC, y + 1, sheet, src, LASTC, src)
        Y["%s%d" % (col, rr)].font = F_B if bold else F_LINK
        Y["%s%d" % (col, rr)].number_format = fmt
        if bold:
            Y["%s%d" % (col, rr)].fill = FILL_TOT
    if bold:
        Y.cell(row=rr, column=1).fill = FILL_TOT
    ann[key + sheet] = rr
    yrow[0] += 1
    return rr


def ysec(text):
    section(Y, yrow[0], text, 6)
    yrow[0] += 1


ysec("Revenue by stream")
for k, lab in [("conc", "Concierge geriatric care"), ("home", "Care at home"), ("cons", "Virtual consultations"),
               ("rehab", "Post-operative rehabilitation"), ("pallh", "Palliative care at home"),
               ("long", "Residence, long-stay"), ("step", "Step-down beds"), ("pall", "Palliative wing"),
               ("pres", "Mortuary preservation"), ("prep", "Preparation and repatriation"),
               ("aut", "Autopsy and pathology"), ("saas", "Software licences")]:
    annual(lab, "Revenue", k)
ysec("Revenue by business line")
for k, lab in [("L1", "Care in the home"), ("L2", "Residential care"), ("L3", "Mortuary and pathology"),
               ("L4", "Software")]:
    annual(lab, "Revenue", k)
r_rev = annual("Revenue", "Revenue", "total", bold=True)
ysec("Costs")
for k, lab in [("DIRECT", "Direct costs"), ("SITE", "Site operating costs"), ("LEASE", "Leases"),
               ("CENTRAL", "Central overhead")]:
    annual(lab, "Costs", k)
r_eb = annual("EBITDA", "P&L", "ebitda", bold=True)
rr = yrow[0]
Y.cell(row=rr, column=1, value="EBITDA margin").font = F_CALC
for y in range(5):
    col = get_column_letter(2 + y)
    Y["%s%d" % (col, rr)] = "=IF(%s%d=0,0,%s%d/%s%d)" % (col, r_rev, col, r_eb, col, r_rev)
    Y["%s%d" % (col, rr)].number_format = PCT
yrow[0] += 1
ysec("Cash")
annual("Leases paid", "Cash Flow", "LPAID")
annual("Capital spend", "Cash Flow", "CAPEX")
annual("Deposits received", "Cash Flow", "dep")
r_net_y = annual("Net cash flow before funding", "Cash Flow", "net", bold=True)
rr = yrow[0]
Y.cell(row=rr, column=1, value="Cumulative cash before funding, year end").font = F_B
for y in range(5):
    col = get_column_letter(2 + y)
    Y["%s%d" % (col, rr)] = "=INDEX('Cash Flow'!$D$%d:$%s$%d,%d)" % (r_cum, LASTC, r_cum, 12 * (y + 1))
    Y["%s%d" % (col, rr)].number_format = NGN
    Y["%s%d" % (col, rr)].font = F_LINK
r_cumy = rr
yrow[0] += 1
rr = yrow[0]
Y.cell(row=rr, column=1, value="Cash balance after the raise, year end").font = F_B
for y in range(5):
    col = get_column_letter(2 + y)
    Y["%s%d" % (col, rr)] = "=INDEX('Cash Flow'!$D$%d:$%s$%d,%d)" % (r_bal, LASTC, r_bal, 12 * (y + 1))
    Y["%s%d" % (col, rr)].number_format = NGN
    Y["%s%d" % (col, rr)].font = F_LINK
yrow[0] += 2
ysec("Year-end operating measures")
for lab, key, fmt in [("Care at home clients", "home", NUM1), ("Concierge members", "conc", NUM1),
                      ("Occupied beds, both residences", "beds", NUM1), ("Bodies held on an average day", "held", NUM1),
                      ("Organisations licensing", "orgs", NUM1)]:
    rr = yrow[0]
    Y.cell(row=rr, column=1, value=lab).font = F_CALC
    src = rows[("Volumes", key)]
    for y in range(5):
        col = get_column_letter(2 + y)
        Y["%s%d" % (col, rr)] = "=INDEX('Volumes'!$D$%d:$%s$%d,%d)" % (src, LASTC, src, 12 * (y + 1))
        Y["%s%d" % (col, rr)].number_format = fmt
        Y["%s%d" % (col, rr)].font = F_LINK
    yrow[0] += 1
Y.freeze_panes = "B6"
name("Y5_EBITDA", Y, "F%d" % r_eb)
name("Y5_REV", Y, "F%d" % r_rev)


# ------------------------------------------------------------ Unit Economics ---
U = wb.create_sheet("Unit Economics")
title(U, "Ilé: unit economics at steady state", "Per unit a month, from the assumptions.")
U.column_dimensions["A"].width = 40
for c in "BCDE":
    U.column_dimensions[c].width = 18
for k, t in enumerate(["Stream", "Revenue a month", "Direct cost", "Contribution", "Margin"]):
    cell = U.cell(row=4, column=1 + k, value=t)
    cell.font = F_SEC
    cell.fill = FILL_SEC
ur = 5
for lab, fee, direct in [("Concierge, per member", "CONC_FEE", "CONC_DIRECT"),
                         ("Care at home, per client", "HOME_ARPU", "HOME_DIRECT"),
                         ("Rehabilitation, per six-week package", "REHAB_FEE", "REHAB_DIRECT"),
                         ("Palliative at home, per patient", "PALLH_FEE", "PALLH_DIRECT"),
                         ("Virtual consultation, each", "CONS_FEE", "CONS_DIRECT"),
                         ("Software, per organisation", "SAAS_FEE", "SAAS_DIRECT")]:
    U.cell(row=ur, column=1, value=lab)
    U.cell(row=ur, column=2, value="=" + fee)
    U.cell(row=ur, column=3, value="=%s*%s" % (fee, direct))
    U.cell(row=ur, column=4, value="=B%d-C%d" % (ur, ur))
    U.cell(row=ur, column=5, value="=IF(B%d=0,0,D%d/B%d)" % (ur, ur, ur))
    ur += 1
U.cell(row=ur, column=1, value="Residence at ceiling, a month")
U.cell(row=ur, column=2, value="=RES_CAP*RES_FEE+STEP_CAP*DAYS_PM*STEP_DAY+PALL_CAP*PALL_FEE")
U.cell(row=ur, column=3, value="=RES_STAFF+(RES_CAP+STEP_CAP+PALL_CAP)*RES_PER_BED+RES_UTIL+RES_OTHER+RES_LEASE_YR/12")
U.cell(row=ur, column=4, value="=B%d-C%d" % (ur, ur))
U.cell(row=ur, column=5, value="=IF(B%d=0,0,D%d/B%d)" % (ur, ur, ur))
ur += 1
U.cell(row=ur, column=1, value="Mortuary at 80 places, a month")
U.cell(row=ur, column=2, value="=MORT_CAP2*MORT_FILL*DAYS_PM*MORT_DAY+MORT_CAP2*MORT_FILL*DAYS_PM/MORT_STAY*MORT_CASE+REPAT_V3*REPAT_FEE+AUT_V3*AUT_FEE")
U.cell(row=ur, column=3, value="=MSTAFF_V2+MPOWER_V2+MORT_OTHER+MORT_LEASE_YR/12+MORT_CAP2*MORT_FILL*DAYS_PM/MORT_STAY*MORT_CASE_COST+REPAT_V3*REPAT_FEE*REPAT_DIRECT+AUT_V3*AUT_FEE*AUT_DIRECT")
U.cell(row=ur, column=4, value="=B%d-C%d" % (ur, ur))
U.cell(row=ur, column=5, value="=IF(B%d=0,0,D%d/B%d)" % (ur, ur, ur))
for rr in range(5, ur + 1):
    for cc in range(2, 5):
        U.cell(row=rr, column=cc).number_format = NGN
    U.cell(row=rr, column=5).number_format = PCT
U.cell(row=ur + 2, column=1, value="Residence and mortuary rows are site EBITDA: they include staffing, running costs and the lease.").font = F_NOTE


# ------------------------------------------------------------------ Market ------
K = wb.create_sheet("Market")
title(K, "Ilé: Lagos market at premium prices", "Bottom up, with checks against mature markets.")
K.column_dimensions["A"].width = 52
K.column_dimensions["B"].width = 20
K.column_dimensions["C"].width = 60
mk = [("Lagos aged 60 and over", "=LAGOS_POP*SHARE_60", "#,##0", ""),
      ("Of them dependent in daily living", "=B4*ADL_DEP", "#,##0", ""),
      ("Deaths a year in Lagos", "=LAGOS_POP*DEATH_RATE", "#,##0", ""),
      ("Care in the home", "=B5*PREMIUM*HOME_ARPU*12", NGN, "Premium dependent elders at the care at home price"),
      ("Residential care", "=B5*NEEDS_24H*PREMIUM*RES_FEE*12", NGN, "Those needing round-the-clock care"),
      ("Palliative care", "=B6*PALL_NEED*PREMIUM*PALL_MONTHS*PALL_FEE", NGN, ""),
      ("Mortuary", "=B6*PRIVATE_MORT*(MORT_STAY*MORT_DAY+MORT_CASE)", NGN, ""),
      ("Lagos premium market a year", "=SUM(B7:B10)", NGN, ""),
      ("Ilé year-five revenue as a share", "=IF(B11=0,0,Y5_REV/B11)", PCT, ""),
      ("Check: residents at the England care-home rate (2.5% of 65+)", "=B4*PREMIUM*0.025", "#,##0", "ONS 2021"),
      ("Check: palliative families at the US hospice rate (51.7%)", "=B6*PREMIUM*0.517", "#,##0", "MedPAC 2023")]
for i, (lab, f, fmt, note) in enumerate(mk):
    K.cell(row=4 + i, column=1, value=lab).font = F_B if "market a year" in lab else F_CALC
    K.cell(row=4 + i, column=2, value=f).number_format = fmt
    K.cell(row=4 + i, column=3, value=note).font = F_NOTE


# ------------------------------------------------------------------ Returns -----
T = wb.create_sheet("Returns")
title(T, "Ilé: the round and returns", "Equity value at exit = year-five EBITDA x multiple + cash on hand.")
T.column_dimensions["A"].width = 40
for c in "BCDE":
    T.column_dimensions[c].width = 18
tr = [("Raise", "=RAISE_USD", USD), ("Stake", "=STAKE", PCT), ("Pre-money", "=PRE_USD", USD),
      ("Post-money", "=POST_USD", USD), ("Year-five EBITDA, naira", "=Y5_EBITDA", NGN),
      ("Cash on hand at month 60, naira", "=INDEX('Cash Flow'!$D$%d:$%s$%d,%d)" % (r_bal, LASTC, r_bal, N), NGN)]
for i, (lab, f, fmt) in enumerate(tr):
    T.cell(row=4 + i, column=1, value=lab)
    T.cell(row=4 + i, column=2, value=f).number_format = fmt
for k, t in enumerate(["Exit multiple", "Equity value ($)", "Investor's stake ($)", "Multiple on raise", "Weak naira equity ($)"]):
    cell = T.cell(row=11, column=1 + k, value=t)
    cell.font = F_SEC
    cell.fill = FILL_SEC
for i, nm in enumerate(["EXIT_LO", "EXIT_MID", "EXIT_HI"]):
    rr = 12 + i
    T.cell(row=rr, column=1, value="=" + nm).number_format = MULT
    T.cell(row=rr, column=2, value="=(A%d*$B$8+$B$9)/FX" % rr).number_format = USD
    T.cell(row=rr, column=3, value="=B%d*STAKE" % rr).number_format = USD
    T.cell(row=rr, column=4, value="=C%d/RAISE_USD" % rr).number_format = MULT
    T.cell(row=rr, column=5, value="=(A%d*$B$8+$B$9)/FX_WEAK" % rr).number_format = USD
T.cell(row=16, column=1, value="Before any later dilution. Software valued at the care multiple, not a software multiple.").font = F_NOTE


# ---------------------------------------------------------------- Scenarios -----
S = wb.create_sheet("Scenarios")
title(S, "Ilé: stress cases", "Recorded from the model. The first four rows can be reproduced with the "
                               "yellow levers on Assumptions; the live result is below.")
S.column_dimensions["A"].width = 52
for c in "BCD":
    S.column_dimensions[c].width = 18
for k, t in enumerate(["Case", "Peak funding (₦)", "Low point", "Year-5 EBITDA (₦)"]):
    cell = S.cell(row=4, column=1 + k, value=t)
    cell.font = F_SEC
    cell.fill = FILL_SEC
order = [0, 1, 2, 6, 3, 4, 5, 7]
for i, idx in enumerate(order):
    nm, kw = M.STRESSES[idx]
    st = M.stressed(**kw)
    S.cell(row=5 + i, column=1, value=nm)
    S.cell(row=5 + i, column=2, value=round(st["peak"])).number_format = NGN
    S.cell(row=5 + i, column=3, value=M.label(st["trough"]["m"]))
    S.cell(row=5 + i, column=4, value=round(st["years"][4]["ebitda"])).number_format = NGN
    for cc in range(2, 5):
        S.cell(row=5 + i, column=cc).font = F_IN
S.cell(row=14, column=1, value="Levers: Residence fill speed 0.5; Mortuary fill speed 0.5; Residence delay 6; Software licences sold 0.").font = F_NOTE
S.cell(row=16, column=1, value="Live, at the current levers").font = F_B
S.cell(row=17, column=1, value="Peak funding need")
S.cell(row=17, column=2, value="=PEAK").number_format = NGN
S.cell(row=18, column=1, value="Raise in naira")
S.cell(row=18, column=2, value="=RAISE_NGN").number_format = NGN
S.cell(row=19, column=1, value="Headroom")
S.cell(row=19, column=2, value="=RAISE_NGN-PEAK").number_format = NGN
S.cell(row=20, column=1, value="Year-five EBITDA")
S.cell(row=20, column=2, value="=Y5_EBITDA").number_format = NGN


# ------------------------------------------------------------------ Cover -------
O = wb.create_sheet("Cover", 0)
title(O, "Ilé: five-year financial model", "October 2026. Confidential. Contact: kemotoks@hotmail.com")
O.column_dimensions["A"].width = 46
O.column_dimensions["B"].width = 22
O.column_dimensions["C"].width = 70
O["A4"] = "Headline"
O["A4"].font = F_B
heads = [("Year-five revenue", "=Y5_REV", NGN), ("Year-five EBITDA", "=Y5_EBITDA", NGN),
         ("Year-five EBITDA margin", "=IF(Y5_REV=0,0,Y5_EBITDA/Y5_REV)", PCT),
         ("Peak funding need", "=PEAK", NGN), ("Month cash is lowest", "=TROUGH_M", "@"),
         ("EBITDA positive from", "=EBITDA_POS", "@"), ("Raise", "=RAISE_USD", USD),
         ("Raise in naira", "=RAISE_NGN", NGN), ("Lowest cash balance after the raise", "=MINBAL", NGN)]
for i, (lab, f, fmt) in enumerate(heads):
    O.cell(row=5 + i, column=1, value=lab)
    cell = O.cell(row=5 + i, column=2, value=f)
    cell.number_format = fmt
    cell.font = F_LINK
O["A16"] = "How the model is built"
O["A16"].font = F_B
guide = [
    ("Assumptions", "Every input, named and blue, with its source. Yellow cells are scenario levers."),
    ("Volumes", "Clients, members, beds, bodies, cases and licences by month."),
    ("Revenue", "Twelve streams in four business lines: care in the home, residential care, mortuary and pathology, software."),
    ("Costs", "Four buckets: direct by stream, site operating, leases, central overhead."),
    ("P&L", "Revenue by line less the four cost buckets gives EBITDA."),
    ("Cash Flow", "EBITDA, leases paid a year ahead, capital spend, deposits, the raise, cumulative cash."),
    ("Annual", "Five-year summary; every figure sums the monthly sheets."),
    ("Unit Economics", "Per member, client, package, residence and mortuary at steady state."),
    ("Market", "Bottom-up Lagos sizing with checks against England and the US."),
    ("Returns", "The round, and the investor's stake at six, eight and ten times EBITDA."),
    ("Scenarios", "Recorded stress cases, and the live result at the current levers."),
]
for i, (s_, t) in enumerate(guide):
    O.cell(row=17 + i, column=1, value=s_).font = F_LINK
    O.cell(row=17 + i, column=2, value=t)
O["A30"] = "Colour code"
O["A30"].font = F_B
O["A31"] = "Blue text: input"
O["A31"].font = F_IN
O["A32"] = "Black text: formula"
O["A33"] = "Green text: link to another sheet"
O["A33"].font = F_LINK
O["A34"] = "Yellow fill: scenario lever"
O["A34"].fill = FILL_KEY
O["A36"] = "Before tax and depreciation. Figures in naira unless marked; ₦1,500 to the dollar. Ilé is a working name."
O["A36"].font = F_NOTE

for ws in wb.worksheets:
    for row_ in ws.iter_rows():
        for cell in row_:
            if cell.font is None or cell.font.name != ARIAL:
                f = cell.font
                cell.font = Font(name=ARIAL, size=f.size or 10, bold=f.bold, italic=f.italic, color=f.color)

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print(OUT)
