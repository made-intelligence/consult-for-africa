"""
Medbury's clinical services proposal, one per specialist partner.

The model is a white label. Medlyfe sells the service under its own brand at
The Lyfe Place, owns the patient, the premises and the capex, and markets it.
The specialist partner delivers the clinical work under a service level
agreement and is paid a share of what is collected. No rent, no capex, no SPV.

The proposal carries the model and the principles only. Fees, service levels
and term are agreed at a meeting, not written into the opening document.

This replaces the three-option frame (lease / rent plus revenue share / SPV)
Medbury sent SS Dental on 31 July 2026. These are Medbury's documents, signed by
Dr Akinware, so they carry Medbury's brand. Built as Word because Medbury edits
and signs before they go.

  python3 scripts/build-medbury-partner-proposals.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
LOGO = ROOT / "docs" / "medbury-logo.png"
OUT = ROOT / "docs" / "medbury-partners"

RED = RGBColor(0xE3, 0x1E, 0x24)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66)
DATE = "9 October 2026"
TBC = "[●]"
BRAND = "Medlyfe"

PARTNERS = [
    {
        "slug": "getslim",
        "name": "GetSlim",
        "service": "metabolic and weight management",
        "why": (
            "Weight and metabolic health sit naturally beside the Fitness and "
            "Health and Wellness pillars, among people already "
            "spending on their health every week."
        ),
        "rooms": (
            "consultation rooms, a body composition and nutrition suite and a "
            "treatment room for injectable therapy"
        ),
        "services": [
            "Medical weight management, including GLP-1 programmes",
            "Bariatric surgery assessment and preparation",
            "Post-operative bariatric follow-up",
            "Metabolic and diabetes risk assessment",
            "Nutrition and dietetic counselling",
            "Body composition analysis",
            "Corporate weight and metabolic health programmes",
        ],
        "surgery": (
            "Bariatric and endoscopic procedures are performed at an accredited "
            "partner hospital and sold by Medlyfe as an all-inclusive package. "
            "Assessment, preparation and follow-up happen at The Lyfe Place."
        ),
        "lead": "a named lead bariatric surgeon",
        "wider": [
            "Medbury would like the partnership to cover GetSlim's whole Lagos "
            "practice, not only the service at The Lyfe Place. GetSlim and "
            "Medlyfe would then work as one weight and metabolic service for "
            "Lagos, with The Lyfe Place as its first shared site.",
            "We would like to agree the shape of that wider partnership "
            "together, including how GetSlim's existing Lagos clinic takes part.",
        ],
        "agenda": "the shape of the wider Lagos partnership",
        "clients_hope": "although we hope Medlyfe at The Lyfe Place becomes GetSlim's destination in Lagos",
    },
    {
        "slug": "oga-dentist",
        "name": "Oga Dentist",
        "service": "dental",
        "why": (
            "Cosmetic and restorative dentistry sit naturally beside the "
            "Aesthetics pillar, and its daily footfall suits a "
            "service built on regular recall."
        ),
        "rooms": (
            "two to three fully equipped dental surgeries with sterilisation "
            "and imaging"
        ),
        "services": [
            "General and preventive dentistry",
            "Cosmetic dentistry and whitening",
            "Restorative dentistry",
            "Orthodontics",
            "Dental implants",
            "Oral surgery",
            "Paediatric dentistry",
            "Corporate dental programmes",
        ],
        "surgery": None,
        "lead": "a named lead dental surgeon",
    },
    {
        "slug": "kp-plastic-surgery",
        "name": "KP Plastic Surgery",
        "service": "aesthetic and plastic surgery",
        "why": (
            "Aesthetics is one of those six pillars, "
            "so the service sits among clients who already value "
            "appearance, wellbeing and longevity."
        ),
        "rooms": (
            "consultation rooms, treatment rooms for non-surgical procedures "
            "and a skin and laser suite"
        ),
        "services": [
            "Aesthetic and plastic surgery consultations",
            "Surgical planning and aftercare",
            "Injectable treatments",
            "Skin and laser treatments",
            "Non-surgical body contouring",
            "Regenerative aesthetics",
        ],
        "surgery": (
            "Surgical procedures are performed at an accredited partner "
            "hospital and sold by Medlyfe as an all-inclusive package. "
            "Consultation, planning and aftercare happen at The Lyfe Place."
        ),
        "lead": "a named lead plastic surgeon",
    },
]


# ---------- page furniture ----------

def _border(par, side: str, sz: int, color: str = "E31E24") -> None:
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color)
    bdr.append(b)


def _shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _styles(doc: Document) -> None:
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(11)
    n.font.color.rgb = INK
    n.paragraph_format.space_after = Pt(8)
    n.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 20), ("Heading 2", 12.5)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RED if name == "Heading 1" else INK
        st.paragraph_format.space_before = Pt(14 if name == "Heading 2" else 0)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def _frame(doc: Document) -> None:
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin, s.bottom_margin = Cm(3.2), Cm(2.2)
    s.different_first_page_header_footer = True

    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.add_run().add_picture(str(LOGO), width=Cm(4.2))
    _border(hp, "bottom", 8)

    for footer in (s.footer, s.first_page_footer):
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Medbury Healthcare Group  |  Clinical Services Proposal  |  Confidential")
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREY


def _runs(p, text: str, **fmt) -> None:
    """Add text to a paragraph with every [●] highlighted."""
    parts = text.split(TBC)
    for i, part in enumerate(parts):
        if part:
            r = p.add_run(part)
            r.bold = fmt.get("bold")
            if "color" in fmt:
                r.font.color.rgb = fmt["color"]
        if i < len(parts) - 1:
            r = p.add_run(TBC)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if fmt.get("bold"):
                r.bold = True


def _text(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _runs(p, text)


def _bullets(doc: Document, items: list[str]) -> None:
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        _runs(p, it)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def _h1(doc: Document, text: str, new_page: bool = True) -> None:
    if new_page:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_heading(text, level=1)
    _border(h, "bottom", 12)


def _table(doc: Document, head: list[str], rows: list[list[str]], widths: list[float]) -> None:
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        _shade(c, "E31E24")
        _runs(c.paragraphs[0], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            lines = val.split("\n")
            _runs(cells[i].paragraphs[0], lines[0], bold=(i == 0 and len(head) > 2))
            for ln in lines[1:]:
                _runs(cells[i].add_paragraph(), ln)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(3)
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


# ---------- the document ----------

def build(pt: dict) -> Path:
    doc = Document()
    _styles(doc)
    _frame(doc)
    name, svc = pt["name"], pt["service"]

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Cm(9))
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CLINICAL SERVICES PROPOSAL")
    r.bold, r.font.size = True, Pt(30)
    _border(p, "bottom", 24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run(
        f"{name} to deliver {svc} services\nat {BRAND}, The Lyfe Place, Lekki Phase 1"
    )
    r.font.size, r.font.color.rgb = Pt(14), GREY
    for _ in range(7):
        doc.add_paragraph()
    for text, colour in (("Prepared By:", RED), ("Medbury Healthcare Group", INK), (DATE, INK)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.bold, r.font.color.rgb = True, colour

    # 1. The proposal
    _h1(doc, "The Proposal")
    _text(doc, (
        f"Medbury Healthcare Group proposes that {BRAND} offers {svc} services to "
        f"its clients at The Lyfe Place, Admiralty Way, Lekki Phase 1, and that "
        f"{name} delivers those services under a service level agreement."
    ))
    _text(doc, (
        f"{BRAND} provides the premises, the fit-out and the equipment, sells the "
        f"service under its brand and markets it. {name} provides the clinicians "
        f"and the clinical standard, and is paid a share of every naira collected "
        f"for the services it delivers. {name} pays no rent and invests no capital."
    ))
    _text(doc, (
        "The Lyfe Place is a premium lifestyle destination built around six "
        "pillars: Fitness, Health and Wellness, Aesthetics, Retail, Community and "
        f"Relaxation. {pt['why']}"
    ))
    if pt.get("wider"):
        doc.add_heading("A partnership across Lagos", level=2)
        for para in pt["wider"]:
            _text(doc, para)

    # 2. What each party brings
    doc.add_heading("What each party brings", level=2)
    _table(doc, [f"{BRAND} (Medbury Healthcare Group)", name], [[
        "Dedicated clinical space at The Lyfe Place\n"
        "Fit-out and all clinical equipment (capital expenditure)\n"
        "Facility registration and licensing\n"
        "Brand, marketing and client acquisition\n"
        "Reception, booking, billing and collections\n"
        "Consumables and clinical supplies\n"
        "Patient records system\n"
        "Corporate clients and Medbury Group referrals",
        f"{pt['lead'][0].upper()}{pt['lead'][1:]} and the clinical team\n"
        "Clinical protocols and standards\n"
        "Clinical governance and quality assurance\n"
        "Professional registration and indemnity for every clinician\n"
        "Clinical input into the service menu and pricing\n"
        f"Training of {BRAND} support staff",
    ]], [8.0, 8.0])

    # 3. The service
    _h1(doc, "The Service")
    _text(doc, f"{BRAND} will fit out {pt['rooms']}, with reception and waiting shared across the Lyfe Place clinical floor. The service menu will include:")
    _bullets(doc, pt["services"])
    if pt["surgery"]:
        _text(doc, pt["surgery"])
    doc.add_heading("How it works for a client", level=2)
    _bullets(doc, [
        f"The client books and pays {BRAND}.",
        f"A {name} clinician sees the client in {BRAND} rooms, to {name}'s protocols.",
        f"Notes are written into {BRAND}'s records system on the day.",
        f"{BRAND} handles follow-up, recall and client care.",
    ])

    # 4. How we would work together
    _h1(doc, "How We Would Work Together")
    _bullets(doc, [
        f"{BRAND} sells the service, sets prices with {name}'s clinical input, and collects all revenue.",
        f"{name} is paid a share of what is collected for the services it delivers.",
        f"{name} employs its own clinicians. {BRAND} employs front-of-house staff.",
        f"An agreement sets service levels on sessions, access, records and client experience, reviewed jointly.",
        f"Clients who come through {BRAND} are {BRAND} clients. {name}'s existing patients remain its own"
        + (f", {pt['clients_hope']}." if pt.get("clients_hope") else "."),
        f"While {name} meets the service levels, {BRAND} will not appoint another provider of {svc} services at The Lyfe Place.",
    ])

    # 5. Next steps
    _h1(doc, "Next Steps", new_page=False)
    _text(doc, (
        f"If the model suits {name}, we would like to meet in the next two weeks "
        "to agree the details together: "
        + (f"{pt['agenda']}, " if pt.get("agenda") else "")
        + "the service menu and prices, the fee, session commitments and "
        "service levels, and the term."
    ))
    _text(doc, "We look forward to your response and to building this service together.")
    for _ in range(3):
        doc.add_paragraph()
    for text, bold in (
        ("Dr Itunu Akinware", True),
        ("Managing Director", False),
        ("Medbury Healthcare Group", False),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(text).bold = bold

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"Medbury Healthcare x {name} - Clinical Services Proposal (Draft).docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for pt in PARTNERS:
        print(build(pt))
