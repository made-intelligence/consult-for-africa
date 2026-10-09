"""
Read a clinic's sales ledger and write aggregate facts as JSON.

    python3 scripts/client-dashboard/analyse_ledger.py <ledger.xlsx> <facts.json> [--year 2026]

Built on Arabella's ledger (one sheet per month, one row per payment, a column
per service line, a SUMMARY sheet), which is the most common shape a small
Nigerian clinic keeps. It finds the header row by its words rather than its
position, so a ledger that moved a column still reads.

What it trusts and what it does not, learned the hard way:
- The MONTH comes from the sheet name, never the date cell. Excel on a
  day-first machine stores 5 January as 1 May; 90 of Arabella's 285 rows were
  swapped and 147 more were text.
- A row needs a payer name. Each month sheet ends in a TOTAL row that holds the
  column sums, and counting it doubles every figure.
- Patients are matched on a normalised name only to count repeat visits. The
  name never leaves this script: it is hashed, and the output holds counts.
"""
from __future__ import annotations

import argparse, collections, datetime, hashlib, json, re

import openpyxl

MONTHS = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY", "AUGUST",
          "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]
NON_SERVICE = {"S/N", "DATE", "PATIENT NAME", "PAYMENT TYPE", "TRANSACTION AMOUNT"}
ORG = re.compile(r"insurance|hmo|ltd|limited|plc|ecowas|clearline|leadway|axa|reliance|allianz|henner|bank|embassy", re.I)


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def channel(pay: str) -> str:
    p = pay.upper()
    if "CASH" in p: return "Cash"
    if "POS" in p: return "Card (POS)"
    if "TRANSFER" in p: return "Bank transfer"
    return "Not recorded"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("ledger"); ap.add_argument("out"); ap.add_argument("--year", type=int, default=datetime.date.today().year)
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.ledger, data_only=True)

    rows, issues = [], collections.Counter()
    sheet_titles = {}
    for mi, m in enumerate(MONTHS, 1):
        if m not in wb.sheetnames: continue
        ws = wb[m]; hdr = None
        first = norm(ws.cell(1, 1).value).upper()
        if first and m not in first and "SALES REPORT" in first: issues["sheet title names the wrong month"] += 1
        for r in ws.iter_rows(values_only=True):
            cells = [norm(c).upper() for c in r]
            if "TRANSACTION AMOUNT" in cells and "DATE" in cells:
                hdr = cells; continue
            if not hdr: continue
            rec = dict(zip(hdr, r))
            name = norm(rec.get("PATIENT NAME"))
            if not name or name.upper().startswith("TOTAL"): continue
            try: amt = float(rec.get("TRANSACTION AMOUNT") or 0)
            except (TypeError, ValueError): issues["amount not a number"] += 1; continue
            d = rec.get("DATE"); day = None
            if isinstance(d, datetime.datetime):
                if d.month == mi: day = d.day
                elif d.day == mi: day = d.month; issues["date with day and month swapped"] += 1
                else: issues["date outside its month"] += 1
            elif d: issues["date typed as text"] += 1
            else: issues["no date"] += 1
            lines = {}
            for k, v in rec.items():
                if k and k not in NON_SERVICE and isinstance(v, (int, float)) and v:
                    label = re.sub(r"\s*\(.*\)$", "", k).strip().title()
                    lines[label] = lines.get(label, 0) + float(v)
            if abs(sum(lines.values()) - amt) > 1: issues["service lines do not add up to the amount"] += 1
            pay = norm(rec.get("PAYMENT TYPE"))
            rows.append(dict(m=mi, day=day, pid=hashlib.sha1(name.lower().encode()).hexdigest()[:12],
                             org=bool(ORG.search(name)), amt=amt, lines=lines, ch=channel(pay),
                             bank=(re.search(r"FIRST BANK|GTB|ACCESS|ZENITH|UBA|STANBIC|FIDELITY", pay.upper()) or [None])[0]))

    months = sorted({r["m"] for r in rows})
    by = {m: [r for r in rows if r["m"] == m] for m in months}
    monthly = [dict(month=m, label=MONTHS[m - 1][:3].title(), revenue=sum(r["amt"] for r in by[m]),
                    payments=len(by[m]), patients=len({r["pid"] for r in by[m]}),
                    largest=max((r["amt"] for r in by[m]), default=0)) for m in months]
    svc_total, svc_month = collections.Counter(), {}
    for m in months:
        c = collections.Counter()
        for r in by[m]: c.update(r["lines"])
        svc_month[m] = dict(c); svc_total.update(c)
    first_seen, visits = {}, collections.Counter()
    for r in sorted(rows, key=lambda r: r["m"]):
        first_seen.setdefault(r["pid"], r["m"]); visits[r["pid"]] += 1
    weekday = collections.Counter()
    for r in rows:
        if r["day"]:
            try: weekday[datetime.date(a.year, r["m"], r["day"]).strftime("%a")] += 1
            except ValueError: pass
    amts = sorted(r["amt"] for r in rows)
    q = lambda p: amts[min(len(amts) - 1, int(len(amts) * p))] if amts else 0
    big = sorted(rows, key=lambda r: -r["amt"])[:5]

    summary = {}
    if "SUMMARY" in wb.sheetnames:
        ws = wb["SUMMARY"]; head = None
        for r in ws.iter_rows(values_only=True):
            cells = [norm(c).upper() for c in r]
            if "JANUARY" in cells: head = cells; continue
            if head and cells and cells[0].startswith("TOTAL"):
                for i, h in enumerate(head):
                    if h in MONTHS and isinstance(r[i], (int, float)): summary[MONTHS.index(h) + 1] = float(r[i])
    recon = [dict(month=m, ledger=next(x["revenue"] for x in monthly if x["month"] == m), summary=summary.get(m)) for m in months]

    facts = dict(
        rows=len(rows), months=months, monthly=monthly,
        total=sum(r["amt"] for r in rows),
        service_total=dict(svc_total.most_common()), service_month=svc_month,
        channel_amount={k: sum(r["amt"] for r in rows if r["ch"] == k) for k in {r["ch"] for r in rows}},
        bank_amount={k: sum(r["amt"] for r in rows if r["bank"] == k) for k in {r["bank"] for r in rows if r["bank"]}},
        org_payer=dict(rows=sum(r["org"] for r in rows), amount=sum(r["amt"] for r in rows if r["org"])),
        patients=len(visits), one_visit=sum(1 for v in visits.values() if v == 1),
        three_plus=sum(1 for v in visits.values() if v >= 3),
        new_by_month={m: sum(1 for v in first_seen.values() if v == m) for m in months},
        weekday=dict(weekday), dated_rows=sum(weekday.values()),
        ticket=dict(p25=q(.25), median=q(.5), p75=q(.75), p90=q(.9)),
        largest=[dict(month=r["m"], amount=r["amt"], lines=r["lines"]) for r in big],
        issues=dict(issues), reconciliation=recon,
    )
    json.dump(facts, open(a.out, "w"), indent=1, default=str)
    print(f"{len(rows)} payments, {len(visits)} patients, {facts['total']:,.0f} across {len(months)} months -> {a.out}")


if __name__ == "__main__":
    main()
